"""Pruebas locales (sin red) de georreferenciar_pdv.py.

Ejecutar:  uv run --with pytest pytest -q test_georreferenciar_pdv.py
Los proveedores de estas pruebas son dobles locales: no representan respuestas reales de ninguna API.
"""
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

import georreferenciar_pdv as g

# Las pruebas que leen la base real se saltan si el Excel no está junto al script (no se sube al repositorio)
requiere_base = pytest.mark.skipif(not g.ARCHIVO_ENTRADA.exists(), reason=f"falta {g.ARCHIVO_ENTRADA.name}")


# --- Lectura y diagnóstico --------------------------------------------------
@requiere_base
def test_lee_la_base_real():
    df = g.leer_excel()
    assert len(df) == 38
    assert list(df.columns) == g.COLUMNAS_ORIGINALES
    assert df["SAP"].is_unique


@requiere_base
def test_diagnostico_de_la_base_real():
    df = g.leer_excel()
    diag = [g.diagnosticar_coordenadas(a, b)[0] for a, b in zip(df["LATITUD"], df["LONGIUTD"])]
    assert diag.count("MARCADOR") == 35
    assert diag.count("SIN_GEO") == 3
    assert diag.count("VALIDA") == 0
    assert sum(g.direccion_ausente(x) for x in df["DIRECCION"]) == 3


@pytest.mark.parametrize("lat, lon, esperado", [
    ("4", "4", "MARCADOR"),
    ("sin geo", "sin geo", "SIN_GEO"),
    (None, "-75.1", "VACIA"),
    ("abc", "-75", "NO_NUMERICA"),
    ("95.2", "-75.1", "FUERA_DE_RANGO"),
    ("40.4168", "-3.7038", "FUERA_DE_COLOMBIA"),
    ("4.6097", "-74.0817", "VALIDA"),
    ("6,2442", "-75,5812", "VALIDA"),
])
def test_diagnosticar_coordenadas(lat, lon, esperado):
    assert g.diagnosticar_coordenadas(lat, lon)[0] == esperado


# --- Direcciones -------------------------------------------------------------
@pytest.mark.parametrize("original, normalizada", [
    ("CR 11 10 31", "Carrera 11 # 10-31"),
    ("CL 91 A 70 02", "Calle 91A # 70-02"),
    ("AK 68 90 88", "Avenida Carrera 68 # 90-88"),
    ("CL 37 SUR 40 48", "Calle 37 Sur # 40-48"),
    ("CALLE 33A 72-107", "Calle 33A # 72-107"),
    ("CARRERA 43 # 32A SUR 27", "Carrera 43 # 32A Sur-27"),
    (".CL 1B SUR 32 12.", "Calle 1B Sur # 32-12"),
    ("CR 77BIS 63C 33 ", "Carrera 77 Bis # 63C-33"),
    ("AV 3FN 52 46", "Avenida 3F Norte # 52-46"),
    ("CRA 27 No 29- 145", "Carrera 27 # 29-145"),
    ("CL 34 CR 10 ESQ", "Calle 34 con Carrera 10"),
    ("CR 25 CL 50 ESQ", "Carrera 25 con Calle 50"),
    ("TV 64A 26 SUR 50 CC PZ AMERICAS", "Transversal 64A # 26 Sur-50 (CC PZ AMERICAS)"),
    ("AV DE LAS AMERICAS 68A 94", "Avenida De Las Americas # 68A-94"),
    ("DG 51 35 120", "Diagonal 51 # 35-120"),
])
def test_normalizar_direccion(original, normalizada):
    assert g.parsear_direccion(original).normalizada() == normalizada


def test_direccion_original_no_se_modifica():
    d = g.parsear_direccion("CR 59A 79 30")
    assert d.original == "CR 59A 79 30"


@pytest.mark.parametrize("valor", ["SIN DIRECCION", "Sin Direccion", "", None, float("nan")])
def test_direccion_ausente(valor):
    d = g.parsear_direccion(valor)
    assert d.ausente and not d.utilizable and d.normalizada() == ""


def test_direccion_mal_digitada_no_es_utilizable():
    d = g.parsear_direccion("CR 1426 124")
    assert not d.utilizable
    assert any("posible error" in p for p in d.problemas)


def test_coincidencia_de_vias():
    assert g.coincide_via(["Calle 15"], ["Calle 15"])
    assert g.coincide_via(["Calle 15 (Avenida Ferrocarril)"], ["Calle 15"])
    assert g.coincide_via(["Avenida Carrera 68"], g.nombres_via("Avenida Carrera", "68"))
    assert not g.coincide_via(["Calle 15A"], ["Calle 15"])
    assert not g.coincide_via(["Calle 150"], ["Calle 15"])
    assert g.coincide_via(["Avenida Guabinal"], ["Avenida Guavinal"], difuso=True)


# --- Nombres -----------------------------------------------------------------
def test_analizar_nombre_cadena_con_codigo_y_sap():
    n = g.analizar_nombre("0362-EXITO BUENAVISTA BARRANQUILLA 10061681", 10061681)
    assert n.marca == "EXITO" and n.codigo_tienda == "0362" and n.sap_en_nombre
    assert n.sucursal == ["BUENAVISTA"] and n.municipio_en_nombre == "BARRANQUILLA"


def test_analizar_nombre_municipio_en_nombre():
    assert g.analizar_nombre("MAKRO - DOSQUEBRADAS", 1).municipio_en_nombre == "DOSQUEBRADAS"
    assert g.analizar_nombre("0320-EXITO CANAVERAL FLORIDA B.", 1).municipio_en_nombre == "FLORIDABLANCA"


def test_analizar_nombre_olimpica_sao():
    n = g.analizar_nombre("SAO 094 VILLA CAROLINA 10629846", 10629846)
    assert n.marca == "OLIMPICA" and n.codigo_tienda == "094" and n.sucursal == ["VILLA", "CAROLINA"]


def test_evaluar_nombre():
    n = g.analizar_nombre("EXITO FLORESTA", 1)
    c = g.Candidato("t", "Éxito Floresta", "Bogotá", 4.68, -74.07, "shop", nombres=["Éxito Floresta"])
    assert g.evaluar_nombre(n, c) == (True, 1.0)
    otro = g.Candidato("t", "Éxito Express", "Bogotá", 4.6, -74.0, "shop", nombres=["Éxito Express"])
    assert g.evaluar_nombre(n, otro) == (True, 0.0)
    ajeno = g.Candidato("t", "Carulla", "Bogotá", 4.6, -74.0, "shop", nombres=["Carulla"])
    assert g.evaluar_nombre(n, ajeno)[0] is False


def test_negocios_ajenos_a_la_marca_no_cuentan():
    n = g.analizar_nombre("EXITO NIQUIA", 1)
    optica = g.Candidato("t", "Ópticas Visión Éxito", "Bello", 6.33, -75.54, "shop", nombres=["Ópticas Visión Éxito"])
    assert g.evaluar_nombre(n, optica)[0] is False


def test_direccion_coherente_por_via_de_cruce():
    d = g.parsear_direccion("CR 53 98 9")
    c = g.Candidato("t", "Éxito", "", 11.0, -74.8, "shop", via="Calle 98", placa="52-115")
    assert g.direccion_coherente(d, c, misma_via=False, d_via=200)
    lejos = g.Candidato("t", "Éxito", "", 11.0, -74.8, "shop", via="Calle 98", placa="70-10")
    assert not g.direccion_coherente(d, lejos, misma_via=False, d_via=200)


def test_direccion_coherente_por_misma_via():
    d = g.parsear_direccion("CL 9 SUR 29D 76")
    c = g.Candidato("t", "Supermú", "", 6.19, -75.56, "shop", via="Calle 9 Sur", placa="29D-80")
    assert g.direccion_coherente(d, c, misma_via=True, d_via=10)


def test_centro_comercial():
    cc = g.Candidato("t", "Centro Comercial Salitre Plaza", "", 4.65, -74.1, "shop", tipo="mall",
                     nombres=["Centro Comercial Salitre Plaza"])
    assert g.es_centro_comercial(cc, ["PLAZA", "SALITRE"])
    assert not g.es_centro_comercial(cc, ["AMERICAS"])


def test_numeros_vecinos():
    assert g.numeros_vecinos("90") == ["91", "89", "92", "88"]
    assert g.numeros_vecinos("23 Sur") == ["24 Sur", "22 Sur", "25 Sur", "21 Sur"]
    assert g.numeros_vecinos(None) == []


# --- Geometría y validación ---------------------------------------------------
def test_distancias_e_intersecciones():
    assert g.distancia_m(4.0, -74.0, 4.0, -74.0) == 0
    assert g.distancia_m(4.0, -74.0, 5.0, -74.0) == pytest.approx(111_195, abs=5)
    calle = [[(4.600, -74.010), (4.600, -74.000)]]
    carrera = [[(4.595, -74.005), (4.605, -74.005)]]
    cruces = g.intersecciones_lineas(calle, carrera)
    assert len(cruces) == 1
    assert abs(cruces[0][0] - 4.600) < 1e-6 and abs(cruces[0][1] + 74.005) < 1e-6
    assert g.distancia_a_lineas(4.6009, -74.005, calle) == pytest.approx(99.5, abs=1)


def test_punto_en_poligono():
    cuadro = {"type": "Polygon", "coordinates": [[[-74, 4], [-73, 4], [-73, 5], [-74, 5], [-74, 4]]]}
    assert g.punto_en_poligono(4.5, -73.5, cuadro) is True
    assert g.punto_en_poligono(5.5, -73.5, cuadro) is False


def test_validar_coordenadas():
    muni = g.Municipio("X", "X", "", 4.6, -74.08, (4.4, 4.8, -74.3, -73.9), None)
    assert g.validar_coordenadas(4.65, -74.05, "co", muni) == []
    assert "Coordenadas no numéricas" in g.validar_coordenadas("4", "x")
    assert "Fuera de Colombia" in g.validar_coordenadas(4, 4)
    assert any("centroide" in f for f in g.validar_coordenadas(4.6001, -74.0801, "co", muni))


# --- Decisión con un proveedor local -------------------------------------------
class ProveedorLocal:
    """Doble de pruebas: devuelve datos fijos definidos en la prueba."""

    def __init__(self, lugares, punto=None):
        self.lugares, self.punto = lugares, punto
        self.muni = g.Municipio("PEREIRA", "Pereira", "Risaralda", 4.80, -75.70,
                                (4.70, 4.90, -75.80, -75.60), None)

    def municipio(self, codigo):
        return self.muni if codigo == "PEREIRA" else None

    def ubicar_direccion(self, d, muni):
        return self.punto

    def buscar_lugares(self, termino, muni, limite=15):
        return list(self.lugares)

    def buscar_libre(self, termino):
        return []


FILA = {"SAP": 1, "PDV": "EXITO CUBA", "CIUDAD": "PEREIRA", "DIRECCION": "CL 72 23 53",
        "LATITUD": "4", "LONGIUTD": "4"}
CALLE_72 = [[(4.8050, -75.7450), (4.8050, -75.7350)]]


def _cand(nombre, lat, lon, pid):
    return g.Candidato("Prueba", nombre, f"{nombre}, Pereira", lat, lon, "shop", place_id=pid,
                       nombres=[nombre], pais="co")


def test_sede_identificada_con_direccion():
    punto = g.PuntoDireccion(4.8050, -75.7400, "CRUCE_APROX", "cruce", "Prueba", CALLE_72)
    prov = ProveedorLocal([_cand("Éxito Cuba", 4.8053, -75.7402, "n/1")], punto)
    r = g.procesar_registro(FILA, prov)
    assert r["ESTADO_GEORREFERENCIACION"] == "GEORREFERENCIADO"
    assert r["LATITUD_VALIDADA"] == pytest.approx(4.8053)
    assert r["PLACE_ID"] == "n/1"


def test_sede_lejos_de_la_direccion_queda_en_revisar():
    punto = g.PuntoDireccion(4.8050, -75.7400, "CRUCE_APROX", "cruce", "Prueba", CALLE_72)
    prov = ProveedorLocal([_cand("Éxito Cuba", 4.8200, -75.7000, "n/1")], punto)
    r = g.procesar_registro(FILA, prov)
    assert r["ESTADO_GEORREFERENCIACION"] == "REVISAR"
    assert r["LATITUD_VALIDADA"] is None


def test_varias_sedes_sin_direccion_queda_en_revisar():
    fila = dict(FILA, PDV="EXITO", DIRECCION="SIN DIRECCION")
    prov = ProveedorLocal([_cand("Éxito", 4.81, -75.71, "n/1"), _cand("Éxito", 4.85, -75.65, "n/2")])
    r = g.procesar_registro(fila, prov)
    assert r["ESTADO_GEORREFERENCIACION"] == "REVISAR"
    assert "2 sedes posibles" in r["OBSERVACIONES"]


def test_solo_cruce_de_vias_es_aproximado():
    punto = g.PuntoDireccion(4.8050, -75.7400, "CRUCE_APROX", "Cruce de Calle 72 con Carrera 23", "Prueba", CALLE_72)
    r = g.procesar_registro(FILA, ProveedorLocal([], punto))
    assert r["ESTADO_GEORREFERENCIACION"] == "REVISAR"
    assert r["TIPO_COINCIDENCIA"] == "UBICACION_APROXIMADA"
    assert "APROXIMADA" in r["OBSERVACIONES"] and r["LATITUD_VALIDADA"] is None


def test_sin_resultado():
    r = g.procesar_registro(FILA, ProveedorLocal([]))
    assert r["ESTADO_GEORREFERENCIACION"] == "SIN_RESULTADO"


def test_respuesta_google_se_convierte_en_candidato():
    # Estructura de Places API (New) Text Search según su documentación; valores ficticios.
    datos = {"places": [{"id": "ChIJ-ejemplo", "displayName": {"text": "Éxito Floresta"},
                         "formattedAddress": "Ak. 68 #90-88, Bogotá, Colombia",
                         "location": {"latitude": 4.68, "longitude": -74.07},
                         "types": ["supermarket", "store"],
                         "addressComponents": [{"longText": "Colombia", "shortText": "CO", "types": ["country"]}]}]}
    c = g.GooglePlaces.candidatos_desde_respuesta(datos, "q")[0]
    assert (c.nombre, c.place_id, c.pais, c.categoria) == ("Éxito Floresta", "ChIJ-ejemplo", "co", "shop")


# --- Correcciones manuales -----------------------------------------------------
@pytest.mark.parametrize("texto, esperado", [
    ("4.823696, -75.684668", (4.823696, -75.684668, "texto")),
    ("4,823696; -75,684668", (4.823696, -75.684668, "texto")),
    ("https://www.google.com/maps/place/Makro/@4.8236,-75.6846,17z/data=!3m1!4b1!4m6!3m5!1s0x0:0x1!8m2"
     "!3d4.8236959!4d-75.6846679", (4.8236959, -75.6846679, "pin")),
    ("https://www.google.com/maps/@6.2442,-75.5812,15z", (6.2442, -75.5812, "vista")),
    ("https://www.google.com/maps/search/?api=1&query=3.45,-76.53", (3.45, -76.53, "consulta")),
    ("no sé", None),
    (None, None),
])
def test_extraer_coordenadas(texto, esperado):
    assert g.extraer_coordenadas(texto, resolver_enlaces=False) == esperado


def test_aplicar_correccion_valida_y_fuera_de_municipio():
    prov = ProveedorLocal([])
    base = {"ESTADO_GEORREFERENCIACION": "REVISAR", "OBSERVACIONES": "Candidato propuesto: 4.800000, -75.700000.",
            "LATITUD_VALIDADA": None, "LONGITUD_VALIDADA": None}
    ok = g.aplicar_correccion(base, {"ENLACE_O_COORDENADAS": "4.8010, -75.7010", "FUENTE": "Google Maps"}, FILA, prov)
    assert ok["ESTADO_GEORREFERENCIACION"] == "GEORREFERENCIADO"
    assert ok["TIPO_COINCIDENCIA"] == "VALIDACION_MANUAL" and ok["LATITUD_VALIDADA"] == 4.801
    assert "Google Maps" in ok["FUENTE_COORDENADAS"] and "m de la ubicación automática" in ok["OBSERVACIONES"]
    fuera = g.aplicar_correccion(base, {"ENLACE_O_COORDENADAS": "6.25, -75.56"}, FILA, prov)
    assert fuera["ESTADO_GEORREFERENCIACION"] == "REVISAR" and fuera["LATITUD_VALIDADA"] is None
    assert "no está dentro de Pereira" in fuera["OBSERVACIONES"]
    ilegible = g.aplicar_correccion(base, {"ENLACE_O_COORDENADAS": "xx"}, FILA, prov)
    assert ilegible["ESTADO_GEORREFERENCIACION"] == "REVISAR"


@requiere_base
def test_plantilla_conserva_lo_escrito(tmp_path):
    df = g.leer_excel()
    resultados = {str(s): {"ESTADO_GEORREFERENCIACION": "REVISAR", "OBSERVACIONES": "x"} for s in df["SAP"]}
    resultados[str(df["SAP"].iloc[0])] = {"ESTADO_GEORREFERENCIACION": "GEORREFERENCIADO", "OBSERVACIONES": "ok"}
    salida = g.construir_salida(df, resultados)
    ruta = tmp_path / "corr.xlsx"
    assert g.generar_plantilla_correcciones(salida, ruta) == 37
    plantilla = pd.read_excel(ruta, dtype=object)
    plantilla.loc[0, "ENLACE_O_COORDENADAS"] = "4.8, -75.7"
    plantilla.to_excel(ruta, index=False)
    g.generar_plantilla_correcciones(salida, ruta)
    corr = g.cargar_correcciones(ruta)
    assert list(corr) == [str(plantilla.loc[0, "SAP"])]
    assert corr[str(plantilla.loc[0, "SAP"])]["ENLACE_O_COORDENADAS"] == "4.8, -75.7"


# --- Exportación --------------------------------------------------------------
@requiere_base
def test_exportar_conserva_registros_y_resumen(tmp_path):
    df = g.leer_excel()
    sap0 = str(df["SAP"].iloc[0])
    resultados = {sap0: {"ESTADO_GEORREFERENCIACION": "GEORREFERENCIADO", "LATITUD_VALIDADA": 4.4,
                         "LONGITUD_VALIDADA": -75.2, "OBSERVACIONES": "ok"}}
    salida = g.construir_salida(df, resultados)
    ruta = tmp_path / "salida.xlsx"
    g.exportar_excel(salida, ruta, "prueba", "prueba")
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Hoja4", "Resumen"]
    leido = pd.read_excel(ruta, sheet_name="Hoja4", dtype=object)
    assert len(leido) == 38 and leido["SAP"].tolist() == df["SAP"].tolist()
    assert list(leido.columns) == g.COLUMNAS_ORIGINALES + g.COLUMNAS_NUEVAS
    assert leido.loc[0, "LATITUD"] == 4.4 and leido.loc[0, "LATITUD_ORIGINAL"] == "4"
    assert leido.loc[1, "LATITUD"] == "4"   # no georreferenciado: original intacto
    resumen = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=2, max_row=9)}
    assert resumen["Total de puntos de venta"] == 38
    assert resumen["Registros georreferenciados"] == 1
    assert resumen["Registros no procesados"] == 37
