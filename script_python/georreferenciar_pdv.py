# /// script
# requires-python = ">=3.10"
# dependencies = ["pandas>=2.0", "openpyxl>=3.1", "requests>=2.31"]
# ///
"""Georreferenciación de los PDV de `PDV multimarca AU.xlsx`.

Proveedores:
  * osm    -> OpenStreetMap / Nominatim (gratuito, máx. 1 consulta por segundo).
  * google -> Google Places API (New) Text Search + Geocoding API.
              Requiere la variable de entorno GOOGLE_MAPS_API_KEY (tiene costo).
  * auto   -> google si existe GOOGLE_MAPS_API_KEY; si no, osm.

Los límites municipales (para validar que el punto cae en el municipio) se
obtienen siempre de Nominatim.

Uso:
  uv run georreferenciar_pdv.py --diagnostico        # sin red: revisa la base
  uv run georreferenciar_pdv.py --prueba             # 5 registros -> archivo _PRUEBA
  uv run georreferenciar_pdv.py                      # 38 registros -> archivo final
"""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import math
import os
import re
import sys
import time
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERSION_ALGORITMO = "1.10"

BASE = Path(__file__).resolve().parent
ARCHIVO_ENTRADA = BASE / "PDV multimarca AU.xlsx"
HOJA_ENTRADA = "Hoja4"
ARCHIVO_SALIDA = BASE / "PDV_multimarca_AU_georreferenciado.xlsx"
ARCHIVO_SALIDA_PRUEBA = BASE / "PDV_multimarca_AU_georreferenciado_PRUEBA.xlsx"
ARCHIVO_CORRECCIONES = BASE / "correcciones_manuales.xlsx"
DIR_TRABAJO = BASE / "georef_trabajo"

COLUMNAS_ORIGINALES = ["SAP", "PDV", "REGIONAL", "CIUDAD", "DIRECCION", "CANAL",
                       "RUTA", "fech", "LATITUD", "LONGIUTD"]
COLUMNAS_NUEVAS = ["LATITUD_ORIGINAL", "LONGITUD_ORIGINAL", "NOMBRE_ENCONTRADO",
                   "DIRECCION_ENCONTRADA", "CIUDAD_ENCONTRADA", "LATITUD_VALIDADA",
                   "LONGITUD_VALIDADA", "ESTADO_GEORREFERENCIACION", "FUENTE_COORDENADAS",
                   "PLACE_ID", "TIPO_COINCIDENCIA", "OBSERVACIONES"]

ESTADOS = ["GEORREFERENCIADO", "REVISAR", "SIN_RESULTADO", "ERROR"]
# Colombia continental + San Andrés: lat_min, lat_max, lon_min, lon_max
COLOMBIA_BBOX = (-4.3, 13.6, -82.0, -66.8)

# Umbrales de validación (metros)
MAX_DIST_VIA = 250        # POI a la vía principal, si el nombre de la sede coincide completo
MAX_DIST_VIA_ESTRICTA = 120   # POI a la vía principal, en los demás casos
MAX_DIST_SOLO_VIA = 60    # POI a la vía cuando no se pudo ubicar el cruce
MAX_DIST_CRUCE = 450      # POI al cruce de vías de la dirección
RADIO_MISMA_SEDE = 150    # candidatos más cercanos que esto se tratan como el mismo lugar
MIN_DIST_CENTRO = 25      # un punto prácticamente igual al centroide municipal es sospechoso
MAX_DIST_PLACA = 150     # POI al punto de una dirección geocodificada por placa (Google)
MAX_DIST_CC = 200        # local de la marca a su centro comercial
MAX_DIST_CONFLICTO = 1000 # sede identificada por nombre pero así de lejos de la dirección -> REVISAR
MAX_DIST_CANDIDATO = 500 # más lejos que esto (sin coincidir el nombre de sede) no se propone como candidata
RADIO_BUSQUEDA_CERCA = 0.006  # grados (~650 m) alrededor de la dirección para buscar la marca

log = logging.getLogger("georef")


# ---------------------------------------------------------------------------
# Texto
# ---------------------------------------------------------------------------
def sin_acentos(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in texto if not unicodedata.combining(c))


def clave(texto) -> str:
    """Forma comparable: sin acentos, minúsculas, solo letras y números."""
    return re.sub(r"[^a-z0-9]", "", sin_acentos(texto).lower())


def tokens(texto) -> list[str]:
    return [t for t in re.split(r"[^A-Z0-9]+", sin_acentos(texto).upper()) if t]


def es_vacio(valor) -> bool:
    return valor is None or (isinstance(valor, float) and math.isnan(valor)) or str(valor).strip() == ""


# ---------------------------------------------------------------------------
# Catálogos
# ---------------------------------------------------------------------------
DEPARTAMENTOS = {
    "IBAGUE": "Tolima", "PEREIRA": "Risaralda", "DOSQUEBRADAS": "Risaralda",
    "LA VIRGINIA": "Risaralda", "BARRANQUILLA": "Atlántico", "SOLEDAD": "Atlántico",
    "PUERTO COLOMBIA": "Atlántico", "MALAMBO": "Atlántico", "BOGOTA": "Bogotá",
    "SOACHA": "Cundinamarca", "FACATATIVA": "Cundinamarca", "ZIPAQUIRA": "Cundinamarca",
    "TUNJA": "Boyacá", "MEDELLIN": "Antioquia", "BELLO": "Antioquia",
    "ENVIGADO": "Antioquia", "SABANETA": "Antioquia", "ITAGUI": "Antioquia",
    "LA ESTRELLA": "Antioquia", "CALI": "Valle del Cauca", "YUMBO": "Valle del Cauca",
    "JAMUNDI": "Valle del Cauca", "BUCARAMANGA": "Santander", "FLORIDABLANCA": "Santander",
    "GIRON": "Santander", "PIEDECUESTA": "Santander", "MANIZALES": "Caldas",
    "VILLAMARIA": "Caldas", "CARTAGENA": "Bolívar",
}
NOMBRE_OFICIAL = {
    "IBAGUE": "Ibagué", "BOGOTA": "Bogotá", "MEDELLIN": "Medellín", "FACATATIVA": "Facatativá",
    "ZIPAQUIRA": "Zipaquirá", "ITAGUI": "Itagüí", "GIRON": "Girón", "JAMUNDI": "Jamundí",
    "VILLAMARIA": "Villamaría",
}
AREA_METROPOLITANA = {
    "PEREIRA": ["DOSQUEBRADAS", "LA VIRGINIA"], "DOSQUEBRADAS": ["PEREIRA"],
    "MEDELLIN": ["BELLO", "ENVIGADO", "ITAGUI", "SABANETA"], "BELLO": ["MEDELLIN"],
    "ENVIGADO": ["MEDELLIN", "SABANETA", "ITAGUI"], "SABANETA": ["ENVIGADO", "MEDELLIN", "LA ESTRELLA"],
    "BARRANQUILLA": ["SOLEDAD", "PUERTO COLOMBIA", "MALAMBO"],
    "BUCARAMANGA": ["FLORIDABLANCA", "GIRON", "PIEDECUESTA"],
    "FLORIDABLANCA": ["BUCARAMANGA", "PIEDECUESTA", "GIRON"],
    "CALI": ["YUMBO", "JAMUNDI"], "BOGOTA": ["SOACHA"], "MANIZALES": ["VILLAMARIA"],
}
# Abreviaturas de municipio que aparecen dentro del nombre del PDV
ALIAS_MUNICIPIO = {("FLORIDA", "B"): "FLORIDABLANCA"}

# codigo, nombre para buscar, alias (secuencias de tokens), ciudades sede de la cadena
MARCAS = [
    ("EXITO", "Éxito", [["EXITO"]], []),
    ("OLIMPICA", "Olímpica", [["OLIMPICA"], ["SAO"]], []),
    ("MAKRO", "Makro", [["MAKRO"]], []),
    ("EURO", "Euro", [["EURO"]], []),
    ("MERCALDAS", "Mercaldas", [["MERCALDAS"]], ["MANIZALES"]),
    ("SUPERMU", "Supermu", [["SUPERMU"], ["SUPER", "MU"]], []),
    ("SUPERINTER", "Superinter", [["SUPERINTER"], ["SUPER", "INTER"]], []),
    ("LAMONTANA", "La Montaña", [["LA", "MONTANA"], ["AGROMERCADOS"], ["LA", "14"]], []),
    ("MASXMENOS", "Mas por Menos", [["MAS", "POR", "MENOS"], ["MAS", "X", "MENOS"], ["MASXMENOS"]], []),
    ("METROPOL", "Metropol", [["METROPOL"]], []),
    ("MEGATIENDAS", "Megatiendas", [["MEGATIENDA"], ["MEGATIENDAS"]], []),
    ("ZAPATOCA", "Zapatoca", [["ZAPATOCA"]], []),
    ("MERCACENTRO", "Mercacentro", [["MERCACENTRO"]], []),
    ("SURTIPLAZA", "Surtiplaza", [["SURTIPLAZA"]], []),
    ("VAQUITA", "La Vaquita", [["VAQUITA"]], []),
]
PALABRAS_GENERICAS = {"SUPERMERCADO", "SUPERMERCADOS", "AUTOSERVICIO", "AUTOSERVICIOS",
                      "MERCADOS", "MERCADO", "SUPERTIENDAS", "SUPERTIENDA", "SAS", "S", "A",
                      "SA", "DE", "LA", "EL", "LOS", "LAS", "DEL", "Y", "N", "NO"}
ABREV_NOMBRE = {"CL": "CALLE", "CLL": "CALLE", "CR": "CARRERA", "CRA": "CARRERA", "KR": "CARRERA",
                "PZ": "PLAZA", "CC": "CENTRO COMERCIAL", "AV": "AVENIDA"}
TIPOS_LUGAR_EXCLUIDOS = {"highway", "place", "boundary", "natural", "waterway", "railway",
                         "route", "landuse"}
# Negocios que comparten la marca pero no son el supermercado (ópticas, droguerías, cajeros...)
NEGOCIOS_AJENOS = {"OPTICA", "OPTICAS", "DROGUERIA", "DROGUERIAS", "FARMACIA", "CAJERO", "ATM",
                   "PARQUEADERO", "ESTACION", "BANCO", "VIAJES", "SEGUROS", "GASOLINERA", "SERVITECA",
                   "CLUB", "VINO", "VINOS", "LICORES", "LICORERA", "AKT", "MOTOS", "BALOTO", "PAC", "EDS"}
# Tipos de Google Places que corresponden a un comercio donde puede estar el PDV
TIPOS_GOOGLE_TIENDA = {"supermarket", "grocery_store", "store", "food_store", "department_store",
                       "hypermarket", "wholesaler", "discount_store", "convenience_store", "shopping_mall",
                       "market", "warehouse_store", "butcher_shop"}
# Palabras de sede demasiado comunes para validarlas solo con la dirección del candidato
SEDE_GENERICA = {"CENTRO", "NORTE", "SUR", "PLAZA", "PARQUE", "CALLE", "CARRERA", "AVENIDA", "EXPRESS"}


# ---------------------------------------------------------------------------
# Lectura y diagnóstico
# ---------------------------------------------------------------------------
def leer_excel(ruta: Path = ARCHIVO_ENTRADA, hoja: str = HOJA_ENTRADA) -> pd.DataFrame:
    df = pd.read_excel(ruta, sheet_name=hoja, dtype=object)
    faltan = [c for c in COLUMNAS_ORIGINALES if c not in df.columns]
    if faltan:
        raise ValueError(f"Faltan columnas en {ruta.name}/{hoja}: {faltan}")
    if df["SAP"].isna().any():
        raise ValueError("Hay registros sin código SAP")
    if df["SAP"].duplicated().any():
        raise ValueError(f"Códigos SAP duplicados: {df.loc[df['SAP'].duplicated(), 'SAP'].tolist()}")
    return df


def diagnosticar_coordenadas(lat, lon) -> tuple[str, str]:
    """Clasifica las coordenadas originales. Solo 'VALIDA' cuenta como válida."""
    if es_vacio(lat) or es_vacio(lon):
        return "VACIA", "Coordenadas vacías"
    if "sin geo" in f"{lat} {lon}".lower():
        return "SIN_GEO", "Coordenadas marcadas como 'sin geo'"
    try:
        la = float(str(lat).replace(",", "."))
        lo = float(str(lon).replace(",", "."))
    except ValueError:
        return "NO_NUMERICA", f"Coordenadas no numéricas ({lat}, {lon})"
    if not (-90 <= la <= 90 and -180 <= lo <= 180):
        return "FUERA_DE_RANGO", f"Coordenadas fuera de rango ({la}, {lo})"
    if la.is_integer() and lo.is_integer():
        extra = "" if en_colombia(la, lo) else "; además el punto cae fuera de Colombia"
        return "MARCADOR", f"Valor marcador {lat}/{lon} pendiente de verificación{extra}"
    if not en_colombia(la, lo):
        return "FUERA_DE_COLOMBIA", f"El punto ({la}, {lo}) está fuera de Colombia"
    return "VALIDA", "Coordenadas originales numéricas y dentro de Colombia (no verificadas)"


def direccion_ausente(direccion) -> bool:
    return es_vacio(direccion) or clave(direccion) in {
        "sindireccion", "sindir", "sd", "na", "nodisponible", "notiene", "0"}


# ---------------------------------------------------------------------------
# Direcciones colombianas
# ---------------------------------------------------------------------------
TIPOS_VIA = {
    "CL": "Calle", "CLL": "Calle", "CALLE": "Calle", "AC": "Avenida Calle",
    "CR": "Carrera", "CRA": "Carrera", "KR": "Carrera", "KRA": "Carrera", "CARRERA": "Carrera",
    "AK": "Avenida Carrera", "DG": "Diagonal", "DIAGONAL": "Diagonal",
    "TV": "Transversal", "TR": "Transversal", "TRANSVERSAL": "Transversal",
    "AV": "Avenida", "AVENIDA": "Avenida",
}
CRUCE_POR_TIPO = {
    "Calle": ["Carrera", "Transversal"], "Avenida Calle": ["Carrera", "Transversal"],
    "Diagonal": ["Carrera", "Transversal"], "Carrera": ["Calle", "Diagonal"],
    "Avenida Carrera": ["Calle", "Diagonal"], "Transversal": ["Calle", "Diagonal"],
    "Avenida": ["Calle", "Carrera"],
}
SENTIDOS = {"SUR": "Sur", "NORTE": "Norte", "ESTE": "Este", "OESTE": "Oeste"}


@dataclass
class Direccion:
    original: str
    ausente: bool = False
    tipo: str | None = None
    nombre_via: str | None = None     # avenidas con nombre (p. ej. "De Las Americas")
    numero: str | None = None         # "91A", "37 Sur", "3F Norte"
    cruce_tipo: str | None = None     # explícito en "CL 34 CR 10 ESQ"
    cruce: str | None = None
    placa: str | None = None
    esquina: bool = False
    complemento: str = ""
    problemas: list[str] = field(default_factory=list)

    @property
    def via_principal(self) -> str | None:
        if not self.tipo:
            return None
        if self.nombre_via:
            return f"{self.tipo} {self.nombre_via}" + (f" {self.numero}" if self.numero else "")
        return f"{self.tipo} {self.numero}" if self.numero else None

    @property
    def utilizable(self) -> bool:
        """Tiene vía principal y cruce: se puede ubicar el tramo."""
        return bool(self.via_principal and self.cruce and not self.ausente
                    and not any("posible error" in p for p in self.problemas))

    def tipos_cruce(self) -> list[str]:
        if self.cruce_tipo:
            return [self.cruce_tipo]
        return CRUCE_POR_TIPO.get(self.tipo or "", ["Calle", "Carrera"])

    def normalizada(self) -> str:
        """Dirección normalizada para consultas (no reemplaza la original)."""
        if self.ausente:
            return ""
        if not self.via_principal:
            return self.original.strip()
        if self.esquina or (self.cruce and not self.placa and self.cruce_tipo):
            texto = f"{self.via_principal} con {self.tipos_cruce()[0]} {self.cruce}"
        elif self.cruce and self.placa:
            texto = f"{self.via_principal} # {self.cruce}-{self.placa}"
        elif self.cruce:
            texto = f"{self.via_principal} # {self.cruce}"
        else:
            texto = self.via_principal
        return texto + (f" ({self.complemento})" if self.complemento else "")


def _leer_numero_via(toks: list[str], i: int) -> tuple[str | None, int]:
    """Lee '91', '91 A', '59A', '77BIS', '37 SUR', '3FN' a partir de toks[i]."""
    if i >= len(toks):
        return None, i
    m = re.fullmatch(r"(\d+)([A-Z]{0,4})", toks[i])
    if not m:
        return None, i
    numero, sufijo = m.group(1), m.group(2)
    i += 1
    letra, bis, sentido = "", False, ""
    if sufijo.endswith("BIS"):
        letra, bis = sufijo[:-3], True
    elif len(sufijo) == 2 and sufijo[1] in "NS":
        letra, sentido = sufijo[0], "Norte" if sufijo[1] == "N" else "Sur"
    else:
        letra = sufijo
    while i < len(toks):
        t = toks[i]
        if not letra and re.fullmatch(r"[A-H]", t):
            letra = t
        elif t == "BIS":
            bis = True
        elif t in SENTIDOS:
            sentido = SENTIDOS[t]
        else:
            break
        i += 1
    texto = numero + letra + (" Bis" if bis else "") + (f" {sentido}" if sentido else "")
    return texto, i


def parsear_direccion(direccion) -> Direccion:
    original = "" if es_vacio(direccion) else str(direccion)
    d = Direccion(original=original)
    if direccion_ausente(direccion):
        d.ausente = True
        d.problemas.append("Dirección ausente")
        return d
    texto = sin_acentos(original).upper()
    texto = re.sub(r"N[°º]|\bNO\.?(?=\s|\d)|\bNRO\.?|#|-|\.|,", " ", texto)
    toks = texto.split()
    if not toks or toks[0] not in TIPOS_VIA:
        d.problemas.append("Tipo de vía no reconocido")
        return d
    d.tipo = TIPOS_VIA[toks[0]]
    i = 1
    if d.tipo == "Avenida" and i < len(toks) and toks[i] in ("CL", "CALLE", "CR", "CRA", "CARRERA", "KR"):
        d.tipo = "Avenida " + TIPOS_VIA[toks[i]]
        i += 1
    if d.tipo == "Avenida" and i < len(toks) and not toks[i][0].isdigit():
        palabras = []
        while i < len(toks) and not toks[i][0].isdigit():
            palabras.append(toks[i])
            i += 1
        d.nombre_via = " ".join(palabras).title()
        # En avenidas con nombre el primer número es el cruce
        d.cruce, i = _leer_numero_via(toks, i)
    else:
        d.numero, i = _leer_numero_via(toks, i)
        if i < len(toks) and toks[i] in TIPOS_VIA:
            d.cruce_tipo = TIPOS_VIA[toks[i]]
            i += 1
        d.cruce, i = _leer_numero_via(toks, i)
    if i < len(toks) and re.fullmatch(r"\d+[A-Z]?", toks[i]):
        d.placa = toks[i]
        i += 1
    resto = toks[i:]
    if any(t in ("ESQ", "ESQUINA") for t in resto):
        d.esquina = True
        resto = [t for t in resto if t not in ("ESQ", "ESQUINA")]
    d.complemento = " ".join(resto)
    if d.cruce_tipo and not d.placa:
        d.esquina = True
    if not d.numero and not d.nombre_via:
        d.problemas.append("Sin número de vía principal")
    if d.numero and len(re.match(r"\d+", d.numero).group()) > 3:
        d.problemas.append(f"Número de vía '{d.numero}' inusual: posible error de digitación")
    if not d.cruce:
        d.problemas.append("Sin número de cruce")
    elif not d.placa and not d.esquina:
        d.problemas.append("Sin número de placa")
    return d


def nombres_via(tipo: str, numero: str | None, nombre: str | None = None) -> list[str]:
    """Nombres con los que puede figurar la vía en OSM/Google."""
    if nombre:
        base = [f"Avenida {nombre}"]
        if numero:
            base.append(f"Avenida {numero}")
        return base
    variantes = [f"{tipo} {numero}"]
    if tipo in ("Calle", "Carrera"):
        variantes.append(f"Avenida {tipo} {numero}")
    if tipo.startswith("Avenida "):
        variantes.append(f"{tipo.split()[1]} {numero}")
    if tipo == "Avenida" and numero:
        variantes.append(f"Avenida {numero.replace(' Norte', 'N').replace(' Sur', 'S')}")
    return variantes


def numeros_vecinos(cruce: str | None) -> list[str]:
    """'90' -> ['91', '89', '92', '88']; conserva el sentido (Sur/Norte)."""
    m = re.match(r"(\d+)", cruce or "")
    if not m:
        return []
    n = int(m.group(1))
    sentido = next((f" {s}" for s in ("Sur", "Norte") if (cruce or "").endswith(s)), "")
    return [f"{n + k}{sentido}" for k in (1, -1, 2, -2) if n + k > 0]


def coincide_via(nombres_osm: list[str], esperados: list[str], difuso: bool = False) -> bool:
    for n in nombres_osm:
        kn = clave(re.sub(r"\(.*?\)", "", str(n)))
        for e in esperados:
            ke = clave(e)
            if kn == ke:
                return True
            if kn.startswith(ke) and not kn[len(ke):len(ke) + 1].isalnum():
                return True
            if difuso and (SequenceMatcher(None, kn, ke).ratio() >= 0.85 or ke in kn):
                return True
    return False


# ---------------------------------------------------------------------------
# Nombre del PDV
# ---------------------------------------------------------------------------
@dataclass
class Nombre:
    original: str
    marca: str | None = None
    marca_busqueda: str | None = None
    alias: list[list[str]] = field(default_factory=list)
    ciudades_marca: list[str] = field(default_factory=list)
    sucursal: list[str] = field(default_factory=list)    # tokens que distinguen la sede
    nucleo: list[str] = field(default_factory=list)      # tokens significativos (sin marca)
    municipio_en_nombre: str | None = None
    codigo_tienda: str | None = None
    sap_en_nombre: bool = False

    def termino_sede(self) -> str:
        partes = [self.marca_busqueda] if self.marca_busqueda else []
        partes += [t.title() for t in self.sucursal]
        return " ".join(partes) or self.original

    def nombre_limpio(self) -> str:
        return " ".join(([self.marca_busqueda] if self.marca_busqueda else []) +
                        [t.title() for t in self.nucleo]) or self.original


def _contiene_secuencia(toks: list[str], seq: list[str]) -> bool:
    n = len(seq)
    return any(toks[i:i + n] == seq for i in range(len(toks) - n + 1))


def _expandir_abreviaturas(toks: list[str]) -> list[str]:
    salida = []
    for t in toks:
        salida.extend(ABREV_NOMBRE.get(t, t).split())
    return salida


def analizar_nombre(pdv: str, sap) -> Nombre:
    nom = Nombre(original=str(pdv))
    texto = sin_acentos(str(pdv)).upper()
    m = re.match(r"\s*(\d{3,4})\s*-", texto) or re.search(r"\bSAO\s+(\d{2,3})\b", texto)
    if m:
        nom.codigo_tienda = m.group(1)
    toks = tokens(texto)
    if str(sap) in toks:
        nom.sap_en_nombre = True
    toks = [t for t in toks if t != str(sap) and t != nom.codigo_tienda]
    for codigo, busqueda, alias, ciudades in MARCAS:
        if any(_contiene_secuencia(toks, a) for a in alias):
            nom.marca, nom.marca_busqueda, nom.alias, nom.ciudades_marca = codigo, busqueda, alias, ciudades
            for a in alias:
                for i in range(len(toks) - len(a) + 1):
                    if toks[i:i + len(a)] == a:
                        toks = toks[:i] + ["\0"] * len(a) + toks[i + len(a):]
            toks = [t for t in toks if t != "\0"]
            break
    for seq, muni in ALIAS_MUNICIPIO.items():
        if _contiene_secuencia(toks, list(seq)):
            nom.municipio_en_nombre = muni
            toks = [t for t in toks if t not in seq]
    for muni in sorted(DEPARTAMENTOS, key=len, reverse=True):
        seq = muni.split()
        if _contiene_secuencia(toks, seq):
            nom.municipio_en_nombre = nom.municipio_en_nombre or muni
            toks = [t for t in toks if t not in seq]
    toks = _expandir_abreviaturas(toks)
    nom.nucleo = [t for t in toks if t not in PALABRAS_GENERICAS]
    nom.sucursal = list(nom.nucleo) if nom.marca else []
    return nom


def evaluar_nombre(nom: Nombre, cand: "Candidato") -> tuple[bool, float]:
    """(coincide la marca o el nombre, fracción de tokens de sede encontrados)."""
    if cand.categoria == "no_tienda":
        return False, 0.0
    toks_nombre = _expandir_abreviaturas(tokens(" ".join(cand.nombres)))
    toks_dir = _expandir_abreviaturas(tokens(cand.direccion))
    if nom.marca:
        marca_ok = any(_contiene_secuencia(toks_nombre, a) for a in nom.alias) and \
            not (NEGOCIOS_AJENOS & set(toks_nombre))
        if not nom.sucursal:
            return marca_ok, 1.0
        # Si la sede tiene palabras distintivas, las genéricas ("Centro", "Plaza") no cuentan
        sede = [t for t in nom.sucursal if t not in SEDE_GENERICA] or nom.sucursal
        puntos = sum(1.0 if t in toks_nombre else
                     0.5 if t in toks_dir and t not in SEDE_GENERICA and not t.isdigit() else 0.0
                     for t in sede)
        return marca_ok, puntos / len(sede)
    if not nom.nucleo:
        return False, 0.0
    frac = sum(t in toks_nombre for t in nom.nucleo) / len(nom.nucleo)
    tienda = cand.categoria in ("shop", "amenity")
    return frac >= 0.99 and tienda, frac


def es_centro_comercial(cand: "Candidato", sucursal: list[str]) -> bool:
    """El candidato es un centro comercial cuyo nombre contiene los tokens de la sede."""
    toks = set(_expandir_abreviaturas(tokens(" ".join(cand.nombres))))
    es_cc = "mall" in cand.tipo or {"CENTRO", "COMERCIAL"} <= toks or "MALL" in toks   # OSM: mall; Google: shopping_mall
    significativos = [t for t in sucursal if t not in ("CENTRO", "COMERCIAL") and t not in SEDE_GENERICA]
    return es_cc and bool(significativos) and all(t in toks for t in significativos)


def _primer_numero(texto) -> int | None:
    m = re.match(r"\s*(\d+)", str(texto or ""))
    return int(m.group(1)) if m else None


def direccion_coherente(d: Direccion, c: "Candidato", misma_via: bool, d_via: float | None) -> bool:
    """El candidato está sobre la vía principal o la de cruce con una placa consistente con la dirección.
    P. ej. 'Carrera 53 # 98-9' y un local en 'Calle 98 # 52-115' describen la misma esquina."""
    if not (d.numero and d.cruce and c.via):
        return False
    n_principal, n_cruce, n_placa = _primer_numero(d.numero), _primer_numero(d.cruce), _primer_numero(c.placa)
    if misma_via and n_placa is not None and n_cruce is not None:
        return abs(n_placa - n_cruce) <= 1
    sobre_cruce = any(coincide_via([c.via], nombres_via(t, d.cruce)) for t in d.tipos_cruce())
    if sobre_cruce and n_placa is not None and n_principal is not None:
        return abs(n_placa - n_principal) <= 2
    return sobre_cruce and d.esquina and d_via is not None and d_via <= MAX_DIST_VIA_ESTRICTA


# ---------------------------------------------------------------------------
# Geometría
# ---------------------------------------------------------------------------
def distancia_m(lat1, lon1, lat2, lon2) -> float:
    r = 6371008.8
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = p2 - p1, math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def _xy(lat, lon, lat0):
    return lon * 111320.0 * math.cos(math.radians(lat0)), lat * 110540.0


def _dist_punto_segmento(p, a, b) -> float:
    (px, py), (ax, ay), (bx, by) = p, a, b
    dx, dy = bx - ax, by - ay
    if dx == dy == 0:
        return math.hypot(px - ax, py - ay)
    t = max(0.0, min(1.0, ((px - ax) * dx + (py - ay) * dy) / (dx * dx + dy * dy)))
    return math.hypot(px - (ax + t * dx), py - (ay + t * dy))


def distancia_a_lineas(lat, lon, lineas: list[list[tuple[float, float]]]) -> float | None:
    """Distancia mínima (m) de un punto a un conjunto de polilíneas [(lat, lon), ...]."""
    mejor = None
    p = _xy(lat, lon, lat)
    for linea in lineas:
        pts = [_xy(a, b, lat) for a, b in linea]
        if len(pts) == 1:
            d = math.hypot(p[0] - pts[0][0], p[1] - pts[0][1])
            mejor = d if mejor is None else min(mejor, d)
        for a, b in zip(pts, pts[1:]):
            d = _dist_punto_segmento(p, a, b)
            mejor = d if mejor is None else min(mejor, d)
    return mejor


def intersecciones_lineas(l1, l2, tolerancia=40.0) -> list[tuple[float, float]]:
    """Todos los cruces entre dos conjuntos de polilíneas (incluye acercamientos < tolerancia),
    agrupando los que distan menos de RADIO_MISMA_SEDE."""
    if not l1 or not l2:
        return []
    lat0 = l1[0][0][0]
    inv = lambda x, y: (y / 110540.0, x / (111320.0 * math.cos(math.radians(lat0))))
    puntos = []
    for a_line in l1:
        a_pts = [_xy(a, b, lat0) for a, b in a_line]
        for b_line in l2:
            b_pts = [_xy(a, b, lat0) for a, b in b_line]
            for p1, p2 in zip(a_pts, a_pts[1:]):
                for q1, q2 in zip(b_pts, b_pts[1:]):
                    d1x, d1y = p2[0] - p1[0], p2[1] - p1[1]
                    d2x, d2y = q2[0] - q1[0], q2[1] - q1[1]
                    den = d1x * d2y - d1y * d2x
                    if den != 0:
                        t = ((q1[0] - p1[0]) * d2y - (q1[1] - p1[1]) * d2x) / den
                        u = ((q1[0] - p1[0]) * d1y - (q1[1] - p1[1]) * d1x) / den
                        if 0 <= t <= 1 and 0 <= u <= 1:
                            puntos.append(inv(p1[0] + t * d1x, p1[1] + t * d1y))
                            continue
                    for pt, s1, s2 in ((p1, q1, q2), (p2, q1, q2), (q1, p1, p2), (q2, p1, p2)):
                        if _dist_punto_segmento(pt, s1, s2) <= tolerancia:
                            puntos.append(inv(*pt))
    unicos: list[tuple[float, float]] = []
    for p in puntos:
        if all(distancia_m(p[0], p[1], u[0], u[1]) > RADIO_MISMA_SEDE for u in unicos):
            unicos.append(p)
    return unicos


def _lineas_geojson(geo: dict | None) -> list[list[tuple[float, float]]]:
    if not geo:
        return []
    t, c = geo.get("type"), geo.get("coordinates")
    if t == "LineString":
        return [[(y, x) for x, y in c]]
    if t == "MultiLineString":
        return [[(y, x) for x, y in l] for l in c]
    if t == "Polygon":
        return [[(y, x) for x, y in c[0]]]
    if t == "Point":
        return [[(c[1], c[0])]]
    return []


def punto_en_poligono(lat, lon, geo: dict | None) -> bool | None:
    if not geo or geo.get("type") not in ("Polygon", "MultiPolygon"):
        return None
    poligonos = [geo["coordinates"]] if geo["type"] == "Polygon" else geo["coordinates"]
    for poli in poligonos:
        anillo, dentro = poli[0], False
        for (x1, y1), (x2, y2) in zip(anillo, anillo[1:] + anillo[:1]):
            if (y1 > lat) != (y2 > lat) and lon < (x2 - x1) * (lat - y1) / (y2 - y1) + x1:
                dentro = not dentro
        if dentro:
            return True
    return False


def en_colombia(lat, lon) -> bool:
    return COLOMBIA_BBOX[0] <= lat <= COLOMBIA_BBOX[1] and COLOMBIA_BBOX[2] <= lon <= COLOMBIA_BBOX[3]


# ---------------------------------------------------------------------------
# Infraestructura HTTP: caché, límite de solicitudes y reintentos
# ---------------------------------------------------------------------------
class ErrorServicio(Exception):
    pass


class Cache:
    """Caché JSON en disco para no repetir consultas idénticas."""

    def __init__(self, ruta: Path):
        self.ruta = ruta
        self.datos = json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}

    @staticmethod
    def llave(*partes) -> str:
        return hashlib.sha1(json.dumps(partes, sort_keys=True, ensure_ascii=False).encode()).hexdigest()

    def obtener(self, llave: str, max_dias: float | None = None):
        e = self.datos.get(llave)
        if e is None:
            return None
        if max_dias is not None and time.time() - e["t"] > max_dias * 86400:
            return None
        return e["r"]

    def guardar(self, llave: str, respuesta, consulta) -> None:
        self.datos[llave] = {"t": time.time(), "q": consulta, "r": respuesta}
        escribir_json(self.ruta, self.datos)


def escribir_json(ruta: Path, datos) -> None:
    tmp = ruta.with_suffix(".tmp")
    tmp.write_text(json.dumps(datos, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    tmp.replace(ruta)


class ClienteHTTP:
    def __init__(self, cache: Cache, intervalo: float, user_agent: str, reintentos: int = 3,
                 max_dias_cache: float | None = None):
        self.cache, self.intervalo, self.reintentos = cache, intervalo, reintentos
        self.max_dias_cache = max_dias_cache
        self.sesion = requests.Session()
        self.sesion.headers["User-Agent"] = user_agent
        self._ultima = 0.0
        self.consultas_reales = 0

    def solicitar(self, metodo: str, url: str, params=None, cuerpo=None, cabeceras=None,
                  ocultar=("key",)):
        visibles = {k: v for k, v in (params or {}).items() if k not in ocultar}
        llave = Cache.llave(metodo, url, visibles, cuerpo, (cabeceras or {}).get("X-Goog-FieldMask"))
        en_cache = self.cache.obtener(llave, self.max_dias_cache)
        if en_cache is not None:
            return en_cache
        ultimo_error = None
        for intento in range(1, self.reintentos + 1):
            espera = self.intervalo - (time.monotonic() - self._ultima)
            if espera > 0:
                time.sleep(espera)
            self._ultima = time.monotonic()
            try:
                r = self.sesion.request(metodo, url, params=params, json=cuerpo,
                                        headers=cabeceras, timeout=30)
                self.consultas_reales += 1
                if r.status_code in (429, 500, 502, 503, 504):
                    pausa = float(r.headers.get("Retry-After", 2 ** intento * 2))
                    ultimo_error = f"HTTP {r.status_code}"
                    log.warning("%s %s -> %s; reintento en %.0fs", metodo, url, ultimo_error, pausa)
                    time.sleep(min(pausa, 60))
                    continue
                if r.status_code >= 400:
                    raise ErrorServicio(f"HTTP {r.status_code}: {r.text[:300]}")
                datos = r.json()
                self.cache.guardar(llave, datos, {"url": url, "params": visibles, "cuerpo": cuerpo})
                return datos
            except (requests.ConnectionError, requests.Timeout, ValueError) as e:
                ultimo_error = f"{type(e).__name__}: {e}"
                log.warning("Fallo de red (%s); intento %d/%d", ultimo_error, intento, self.reintentos)
                time.sleep(2 ** intento)
        raise ErrorServicio(f"Sin respuesta tras {self.reintentos} intentos ({ultimo_error})")


# ---------------------------------------------------------------------------
# Proveedores
# ---------------------------------------------------------------------------
@dataclass
class Municipio:
    codigo: str
    nombre: str
    departamento: str
    lat: float
    lon: float
    bbox: tuple[float, float, float, float]  # lat_min, lat_max, lon_min, lon_max
    geojson: dict | None
    margen: float = 0.01

    def contiene(self, lat, lon) -> bool:
        dentro = punto_en_poligono(lat, lon, self.geojson)
        if dentro is None:
            return self.bbox[0] <= lat <= self.bbox[1] and self.bbox[2] <= lon <= self.bbox[3]
        return dentro

    def viewbox(self) -> str:
        margen = self.margen
        return (f"{self.bbox[2] - margen},{self.bbox[1] + margen},"
                f"{self.bbox[3] + margen},{self.bbox[0] - margen}")


@dataclass
class Candidato:
    fuente: str
    nombre: str
    direccion: str
    lat: float
    lon: float
    categoria: str = ""
    tipo: str = ""
    place_id: str | None = None
    nombres: list[str] = field(default_factory=list)
    via: str = ""
    placa: str = ""
    pais: str = ""
    consulta: str = ""


@dataclass
class PuntoDireccion:
    lat: float | None
    lon: float | None
    precision: str            # EXACTA | ESQUINA | CRUCE_APROX | CRUCE_AMBIGUO | INTERPOLADA | SOLO_VIA
    descripcion: str
    fuente: str
    lineas: list = field(default_factory=list)
    place_id: str | None = None
    alternativas: list = field(default_factory=list)   # todos los cruces encontrados

    def distancia(self, lat, lon) -> float | None:
        puntos = self.alternativas or ([(self.lat, self.lon)] if self.lat is not None else [])
        return min((distancia_m(lat, lon, a, b) for a, b in puntos), default=None)


class Nominatim:
    URL = "https://nominatim.openstreetmap.org/search"
    FUENTE = "OpenStreetMap/Nominatim"

    def __init__(self, cache: Cache):
        ua = os.environ.get("NOMINATIM_USER_AGENT",
                            "PDV-multimarca-AU-georef/1.0 (georreferenciacion interna de 38 PDV)")
        self.http = ClienteHTTP(cache, intervalo=1.1, user_agent=ua)
        self.email = os.environ.get("NOMINATIM_EMAIL")
        self._municipios: dict[str, Municipio | None] = {}

    def _buscar(self, **params) -> list[dict]:
        base = {"format": "jsonv2", "countrycodes": "co", "addressdetails": 1,
                "accept-language": "es", "limit": 10}
        if self.email:
            base["email"] = self.email
        base.update(params)
        return self.http.solicitar("GET", self.URL, params=base, ocultar=("email",))

    def municipio(self, codigo: str) -> Municipio | None:
        if codigo in self._municipios:
            return self._municipios[codigo]
        nombre = NOMBRE_OFICIAL.get(codigo, codigo.title())
        depto = DEPARTAMENTOS.get(codigo)
        q = f"{nombre}, {depto}, Colombia" if depto and depto != nombre else f"{nombre}, Colombia"
        res = self._buscar(q=q, limit=10, polygon_geojson=1, polygon_threshold=0.0005)
        if not any(clave(r.get("name", "")) == clave(nombre) for r in res):
            res = res + self._buscar(q=nombre, limit=10, polygon_geojson=1, polygon_threshold=0.0005)
        # Preferir el límite municipal; si OSM no lo tiene, el perímetro urbano ("Perímetro Urbano X", "X ciudad")
        opciones = []
        for r in res:
            if r.get("category") not in ("boundary", "place") or "geojson" not in r:
                continue
            estado = clave((r.get("address") or {}).get("state", ""))
            if depto and clave(depto) not in estado:
                continue
            nr, kn = clave(r.get("name", "")), clave(nombre)
            nivel = 0 if nr == kn else 1 if nr in (f"perimetrourbano{kn}", f"{kn}ciudad") else None
            if nivel is not None:
                opciones.append((nivel, r.get("place_rank", 99), r))
        elegido = min(opciones, key=lambda o: o[:2])[2] if opciones else None
        if elegido and clave(elegido.get("name", "")) != clave(nombre):
            log.info("Límite de %s: se usa '%s'", nombre, elegido.get("name"))
        m = None
        if elegido:
            bb = [float(x) for x in elegido["boundingbox"]]
            m = Municipio(codigo, nombre, depto or "", float(elegido["lat"]), float(elegido["lon"]),
                          (bb[0], bb[1], bb[2], bb[3]), elegido.get("geojson"))
        else:
            log.warning("No se encontró el límite municipal de %s", q)
        self._municipios[codigo] = m
        return m

    @staticmethod
    def _a_candidato(r: dict, consulta: str) -> Candidato:
        nd = r.get("namedetails") or {}
        et = r.get("extratags") or {}
        nombres = [r.get("name") or ""] + list(nd.values()) + [et.get(k, "") for k in ("brand", "branch", "operator")]
        ad = r.get("address") or {}
        return Candidato(
            fuente=Nominatim.FUENTE, nombre=r.get("name") or "", direccion=r.get("display_name", ""),
            lat=float(r["lat"]), lon=float(r["lon"]), categoria=r.get("category", ""),
            tipo=r.get("type", ""), place_id=f"{r.get('osm_type', '')}/{r.get('osm_id', '')}",
            nombres=[n for n in nombres if n], via=ad.get("road", ""),
            placa=ad.get("house_number", ""), pais=ad.get("country_code", ""), consulta=consulta)

    def buscar_lugares(self, termino: str, muni: Municipio, limite: int = 15) -> list[Candidato]:
        res = self._buscar(q=termino, viewbox=muni.viewbox(), bounded=1, limit=limite,
                           extratags=1, namedetails=1)
        return [self._a_candidato(r, termino) for r in res
                if r.get("category") not in TIPOS_LUGAR_EXCLUIDOS]

    def buscar_libre(self, termino: str) -> list[Candidato]:
        res = self._buscar(q=termino, limit=5, extratags=1, namedetails=1)
        return [self._a_candidato(r, termino) for r in res]

    def _lineas_via(self, nombres: list[str], muni: Municipio, difuso=False) -> list:
        lineas = []
        for n in nombres:
            res = self._buscar(q=f"{n}, {muni.nombre}", viewbox=muni.viewbox(), bounded=1,
                               limit=20, namedetails=1, polygon_geojson=1, polygon_threshold=0.00002)
            for r in res:
                if r.get("category") != "highway":
                    continue
                nombres_osm = [r.get("name", "")] + list((r.get("namedetails") or {}).values())
                if coincide_via(nombres_osm, nombres, difuso) and muni.contiene(float(r["lat"]), float(r["lon"])):
                    lineas.extend(_lineas_geojson(r.get("geojson")))
            if lineas:
                break
        return lineas

    def _interpolar_placa(self, principal, tipo_cruce, numero_cruce, cruce, placa, muni):
        """En la nomenclatura colombiana la placa ('# 10-172') equivale aproximadamente a los metros
        desde el cruce hacia la siguiente vía. Ubica el punto sobre la cuadra entre ambos cruces."""
        metros = _primer_numero(placa)
        n = _primer_numero(numero_cruce)
        if not metros or n is None:
            return None
        sentido = next((f" {s}" for s in ("Sur", "Norte") if str(numero_cruce).endswith(s)), "")
        siguiente = self._lineas_via(nombres_via(tipo_cruce, f"{n + 1}{sentido}"), muni)
        candidatos = [p for p in intersecciones_lineas(principal, siguiente)
                      if 30 <= distancia_m(cruce[0], cruce[1], p[0], p[1]) <= 400]
        if not candidatos:
            return None
        sig = min(candidatos, key=lambda p: distancia_m(cruce[0], cruce[1], p[0], p[1]))
        cuadra = distancia_m(cruce[0], cruce[1], sig[0], sig[1])
        f = min(metros / cuadra, 0.95)
        lat = cruce[0] + f * (sig[0] - cruce[0])
        lon = cruce[1] + f * (sig[1] - cruce[1])
        return lat, lon, f"interpolado {metros} m hacia la {tipo_cruce} {n + 1}{sentido} (cuadra de {cuadra:.0f} m)"

    def ubicar_direccion(self, d: Direccion, muni: Municipio) -> PuntoDireccion | None:
        if not d.via_principal:
            return None
        esperados = nombres_via(d.tipo, d.numero, d.nombre_via)
        # 1) placa exacta registrada en OSM (addr:street + addr:housenumber)
        if d.cruce and d.placa:
            for r in self._buscar(q=f"{d.via_principal} {d.cruce}-{d.placa}, {muni.nombre}",
                                  viewbox=muni.viewbox(), bounded=1, limit=5):
                ad = r.get("address") or {}
                if clave(ad.get("house_number", "")) == clave(f"{d.cruce}{d.placa}") and \
                        coincide_via([ad.get("road", "")], esperados, bool(d.nombre_via)):
                    return PuntoDireccion(float(r["lat"]), float(r["lon"]), "EXACTA",
                                          f"Placa {ad['house_number']} sobre {ad['road']}", self.FUENTE,
                                          place_id=f"{r.get('osm_type')}/{r.get('osm_id')}")
        # 2) geometría de la vía principal y del cruce
        principal = self._lineas_via(esperados, muni, difuso=bool(d.nombre_via))
        if not principal:
            return None
        if d.cruce:
            numeros_cruce = [d.cruce]
            for sentido in ("Norte", "Sur"):   # p. ej. Avenida 3F Norte cruza la Calle 52 Norte
                if d.numero and d.numero.endswith(sentido) and not d.cruce.endswith(sentido):
                    numeros_cruce.insert(0, f"{d.cruce} {sentido}")
            # La vía del cruce puede no llegar físicamente a la principal: probar las vecinas
            vecinos = [] if d.esquina else numeros_vecinos(d.cruce)
            intentos = [(t, n, False) for t in d.tipos_cruce() for n in numeros_cruce] + \
                       [(d.tipos_cruce()[0], n, True) for n in vecinos]
            for tipo_cruce, numero, es_vecino in intentos:
                cruce = self._lineas_via(nombres_via(tipo_cruce, numero), muni)
                cruces = intersecciones_lineas(principal, cruce)
                if cruces:
                    precision = "ESQUINA" if d.esquina else "CRUCE_APROX"
                    desc = f"Cruce de {d.via_principal} con {tipo_cruce} {numero}"
                    if es_vecino:
                        desc += f" (vía vecina; la {tipo_cruce} {d.cruce} no cruza en la cartografía)"
                    if len(cruces) > 1:
                        precision = "CRUCE_AMBIGUO"
                        desc += f" ({len(cruces)} cruces posibles en el municipio)"
                    lat, lon = cruces[0]
                    if len(cruces) == 1 and not es_vecino and not d.esquina:
                        interpolado = self._interpolar_placa(principal, tipo_cruce, numero, cruces[0], d.placa, muni)
                        if interpolado:
                            lat, lon, desc_i = interpolado
                            desc += f"; {desc_i}"
                    return PuntoDireccion(lat, lon, precision, desc, self.FUENTE, principal,
                                          alternativas=[(lat, lon)] + cruces)
        return PuntoDireccion(None, None, "SOLO_VIA", f"Solo se ubicó la vía {d.via_principal}",
                              self.FUENTE, principal)


class GooglePlaces:
    """Google Places API (New) Text Search + Geocoding API. Requiere GOOGLE_MAPS_API_KEY."""
    URL_TEXT = "https://places.googleapis.com/v1/places:searchText"
    URL_GEOCODE = "https://maps.googleapis.com/maps/api/geocode/json"
    CAMPOS = ("places.id,places.displayName,places.formattedAddress,places.location,"
              "places.types,places.addressComponents")
    FUENTE = "Google Places API"

    def __init__(self, cache: Cache, api_key: str, osm: Nominatim):
        # Las condiciones de Google permiten conservar lat/lng hasta 30 días.
        self.http = ClienteHTTP(cache, intervalo=0.2, user_agent="PDV-multimarca-AU-georef/1.0",
                                max_dias_cache=30)
        self.key = api_key
        self.osm = osm

    def municipio(self, codigo: str) -> Municipio | None:
        return self.osm.municipio(codigo)

    @staticmethod
    def candidatos_desde_respuesta(datos: dict, consulta: str) -> list[Candidato]:
        salida = []
        for p in datos.get("places", []):
            comps = p.get("addressComponents", [])
            pais = next((c.get("shortText", "") for c in comps if "country" in c.get("types", [])), "")
            ruta = next((c.get("longText", "") for c in comps if "route" in c.get("types", [])), "")
            num = next((c.get("longText", "") for c in comps if "street_number" in c.get("types", [])), "")
            loc = p.get("location", {})
            if "latitude" not in loc:
                continue
            nombre = (p.get("displayName") or {}).get("text", "")
            tipos = p.get("types", [])
            if {"locality", "political", "route"} & set(tipos):
                categoria = "place"
            elif TIPOS_GOOGLE_TIENDA & set(tipos):
                categoria = "shop"
            else:
                categoria = "no_tienda"   # p. ej. puntos de pago o servicios con el nombre de la marca
            salida.append(Candidato(
                fuente=GooglePlaces.FUENTE, nombre=nombre, direccion=p.get("formattedAddress", ""),
                lat=loc["latitude"], lon=loc["longitude"],
                categoria=categoria,
                tipo=",".join(tipos), place_id=p.get("id"), nombres=[nombre], via=ruta, placa=num,
                pais=pais.lower(), consulta=consulta))
        return salida

    def buscar_lugares(self, termino: str, muni: Municipio, limite: int = 10) -> list[Candidato]:
        cuerpo = {"textQuery": termino, "languageCode": "es", "regionCode": "CO",
                  "pageSize": min(limite, 20),
                  "locationRestriction": {"rectangle": {
                      "low": {"latitude": muni.bbox[0], "longitude": muni.bbox[2]},
                      "high": {"latitude": muni.bbox[1], "longitude": muni.bbox[3]}}}}
        datos = self.http.solicitar("POST", self.URL_TEXT, cuerpo=cuerpo, cabeceras={
            "X-Goog-Api-Key": self.key, "X-Goog-FieldMask": self.CAMPOS})
        return [c for c in self.candidatos_desde_respuesta(datos, termino) if c.categoria != "place"]

    def buscar_libre(self, termino: str) -> list[Candidato]:
        datos = self.http.solicitar("POST", self.URL_TEXT, cuerpo={
            "textQuery": termino, "languageCode": "es", "regionCode": "CO", "pageSize": 5},
            cabeceras={"X-Goog-Api-Key": self.key, "X-Goog-FieldMask": self.CAMPOS})
        return self.candidatos_desde_respuesta(datos, termino)

    def ubicar_direccion(self, d: Direccion, muni: Municipio) -> PuntoDireccion | None:
        if not d.via_principal:
            return None
        datos = self.http.solicitar("GET", self.URL_GEOCODE, params={
            "address": f"{d.normalizada()}, {muni.nombre}, {muni.departamento}, Colombia",
            "components": "country:CO", "language": "es", "key": self.key})
        if datos.get("status") not in ("OK", "ZERO_RESULTS"):
            raise ErrorServicio(f"Geocoding: {datos.get('status')} {datos.get('error_message', '')}")
        for r in datos.get("results", []):
            tipo_loc = r["geometry"]["location_type"]
            if tipo_loc in ("GEOMETRIC_CENTER", "APPROXIMATE") or "locality" in r.get("types", []):
                continue  # centro de vía/ciudad: no sirve como ubicación del PDV
            loc = r["geometry"]["location"]
            if not muni.contiene(loc["lat"], loc["lng"]):
                continue  # Google la ubicó en otro municipio: no es la dirección de este
            if not coincide_via([r.get("formatted_address", "").split(",")[0].replace("Cra.", "Carrera")
                                 .replace("Cl.", "Calle").replace("Ak.", "Avenida Carrera")
                                 .replace("Dg.", "Diagonal").replace("Tv.", "Transversal").split("#")[0]],
                                nombres_via(d.tipo, d.numero, d.nombre_via), bool(d.nombre_via)):
                tipo_loc = "RANGE_INTERPOLATED"  # Google devolvió otra vía: no se considera exacta
            precision = "EXACTA" if tipo_loc == "ROOFTOP" else "INTERPOLADA"
            return PuntoDireccion(loc["lat"], loc["lng"], precision, r.get("formatted_address", ""),
                                  "Google Geocoding API", place_id=r.get("place_id"))
        return None


# ---------------------------------------------------------------------------
# Validación y decisión
# ---------------------------------------------------------------------------
def validar_coordenadas(lat, lon, pais: str = "co", muni: Municipio | None = None) -> list[str]:
    """Lista de fallos; vacía si las coordenadas son aceptables."""
    fallos = []
    try:
        lat, lon = float(lat), float(lon)
    except (TypeError, ValueError):
        return ["Coordenadas no numéricas"]
    if math.isnan(lat) or math.isnan(lon):
        return ["Coordenadas no numéricas"]
    if not (-90 <= lat <= 90 and -180 <= lon <= 180):
        fallos.append("Fuera del rango geográfico")
    if not en_colombia(lat, lon) or (pais and pais.lower() != "co"):
        fallos.append("Fuera de Colombia")
    if muni and distancia_m(lat, lon, muni.lat, muni.lon) < MIN_DIST_CENTRO:
        fallos.append(f"A menos de {MIN_DIST_CENTRO} m del centroide de {muni.nombre}: posible centro de ciudad")
    return fallos


def _agrupar(cands: list[Candidato]) -> list[list[Candidato]]:
    grupos: list[list[Candidato]] = []
    for c in cands:
        for g in grupos:
            if distancia_m(c.lat, c.lon, g[0].lat, g[0].lon) <= RADIO_MISMA_SEDE:
                g.append(c)
                break
        else:
            grupos.append([c])
    return grupos


def codigo_municipio(ciudad) -> str:
    """'Medellín' / 'MEDELLIN' -> 'MEDELLIN' (clave de los catálogos)."""
    k = clave(ciudad).upper()
    return next((m for m in DEPARTAMENTOS if clave(m).upper() == k), str(ciudad).strip().upper())


def _fmt(lat, lon) -> str:
    return f"{lat:.6f}, {lon:.6f}"


def procesar_registro(fila: dict, prov) -> dict:
    """Georreferencia un PDV y devuelve las columnas nuevas + detalle de auditoría."""
    obs: list[str] = []
    sap, pdv, ciudad = fila["SAP"], str(fila["PDV"]), str(fila["CIUDAD"]).strip()
    diag, diag_txt = diagnosticar_coordenadas(fila["LATITUD"], fila["LONGIUTD"])
    obs.append(f"Original: {diag_txt}.")
    nom = analizar_nombre(pdv, sap)
    d = parsear_direccion(fila["DIRECCION"])
    if nom.codigo_tienda:
        obs.append(f"Código de tienda en el nombre: {nom.codigo_tienda}.")
    if d.problemas:
        obs.append("Dirección: " + "; ".join(d.problemas) + ".")
    consultas: list[str] = []
    res = {"NOMBRE_ENCONTRADO": "", "DIRECCION_ENCONTRADA": "", "CIUDAD_ENCONTRADA": "",
           "LATITUD_VALIDADA": None, "LONGITUD_VALIDADA": None, "FUENTE_COORDENADAS": "",
           "PLACE_ID": "", "TIPO_COINCIDENCIA": "", "DIAGNOSTICO_ORIGINAL": diag}

    ciudad_cod = codigo_municipio(ciudad)
    # Municipios a revisar: registrado, el indicado en el nombre, sedes de la cadena, área metropolitana
    principales = [ciudad_cod]
    if nom.municipio_en_nombre and nom.municipio_en_nombre not in principales:
        principales.append(nom.municipio_en_nombre)
    secundarios = [m for m in AREA_METROPOLITANA.get(ciudad_cod, []) + nom.ciudades_marca
                   if m not in principales]

    def evaluar_en(codigos: list[str]):
        todos, puntos = [], {}
        for cod in codigos:
            muni = prov.municipio(cod)
            if not muni:
                continue
            if d.utilizable or (d.via_principal and d.esquina):
                punto = prov.ubicar_direccion(d, muni)
                consultas.append(f"[P2] dirección normalizada '{d.normalizada()}' en {muni.nombre}: "
                                 f"{punto.precision if punto else 'sin resultado'}")
                if punto:
                    puntos[cod] = (muni, punto)
            terminos = [nom.termino_sede(), nom.marca_busqueda] if nom.marca else [nom.nombre_limpio()]
            for t in dict.fromkeys(x for x in terminos if x):
                cands = prov.buscar_lugares(t, muni, limite=40 if t == nom.marca_busqueda else 15)
                consultas.append(f"[P3] '{t}' en {muni.nombre}: {len(cands)} lugares")
                for c in cands:
                    if muni.contiene(c.lat, c.lon):
                        todos.append((muni, c))
        return todos, puntos

    # Prioridad 1: nombre + dirección + ciudad en una sola consulta
    p1 = ", ".join(x for x in [nom.nombre_limpio(), d.normalizada(), ciudad.title(), "Colombia"] if x)
    cands_p1 = prov.buscar_libre(p1)
    consultas.append(f"[P1] '{p1}': {len(cands_p1)} resultados")

    def hay_sede_probable(lugares_):
        for _, c in lugares_:
            ok_nombre, frac = evaluar_nombre(nom, c)
            if ok_nombre and frac >= (0.5 if nom.sucursal else 1.0):
                return True
        return False

    lugares, puntos = evaluar_en(principales)
    evaluados_muni = list(principales)
    if not hay_sede_probable(lugares) and secundarios:
        mas_lugares, mas_puntos = evaluar_en(secundarios)
        lugares += mas_lugares
        evaluados_muni += secundarios
        for k, v in mas_puntos.items():
            puntos.setdefault(k, v)
    # P1 aporta candidatos si caen en algún municipio evaluado
    for c in cands_p1:
        if c.categoria in TIPOS_LUGAR_EXCLUIDOS:
            continue
        for cod in evaluados_muni:
            muni = prov.municipio(cod)
            if muni and muni.contiene(c.lat, c.lon):
                lugares.append((muni, c))
                break

    # Centros comerciales mencionados en el nombre o la dirección (p. ej. "Plaza Salitre", "Cosmocentro")
    centros: list[tuple[Municipio, Candidato]] = []
    termino_cc = " ".join(t.title() for t in nom.sucursal if t not in ("CENTRO", "COMERCIAL"))
    if any(t not in SEDE_GENERICA and not t.isdigit() for t in nom.sucursal):
        for cod in principales:
            muni = prov.municipio(cod)
            if not muni:
                continue
            for t in (f"Centro Comercial {termino_cc}", termino_cc):
                encontrados = [c for c in prov.buscar_lugares(t, muni, 10)
                               if es_centro_comercial(c, nom.sucursal) and muni.contiene(c.lat, c.lon)]
                consultas.append(f"[P3] centro comercial '{t}' en {muni.nombre}: {len(encontrados)}")
                if encontrados:
                    centros += [(muni, c) for c in encontrados]
                    break
    # Búsqueda de la marca alrededor de la dirección y de los centros comerciales: en ciudades con
    # muchos locales de la misma cadena la búsqueda municipal no los devuelve todos
    referencias = [(m, p.lat, p.lon) for cod, (m, p) in puntos.items()
                   if p.lat is not None and cod in principales]
    referencias += [(m, c.lat, c.lon) for m, c in centros[:2]]
    termino_cerca = nom.marca_busqueda or nom.nombre_limpio()
    for muni, lat, lon in referencias[:3]:
        caja = Municipio(muni.codigo, muni.nombre, muni.departamento, muni.lat, muni.lon,
                         (lat - RADIO_BUSQUEDA_CERCA, lat + RADIO_BUSQUEDA_CERCA,
                          lon - RADIO_BUSQUEDA_CERCA, lon + RADIO_BUSQUEDA_CERCA), None, margen=0)
        cerca = prov.buscar_lugares(termino_cerca, caja, 20)
        consultas.append(f"[P3] '{termino_cerca}' alrededor de {_fmt(lat, lon)}: {len(cerca)} lugares")
        lugares += [(muni, c) for c in cerca if muni.contiene(c.lat, c.lon)]

    # Deduplicar por id de lugar
    vistos, unicos = set(), []
    for muni, c in lugares:
        if c.place_id not in vistos:
            vistos.add(c.place_id)
            unicos.append((muni, c))

    # Locales de la marca por municipio (una marca sin sufijo de sede y un único local identifica la sede)
    locales_marca: dict[str, int] = {}
    for muni, c in unicos:
        if evaluar_nombre(nom, c)[0]:
            locales_marca.setdefault(muni.codigo, []).append(c)
    locales_marca = {k: len(_agrupar(v)) for k, v in locales_marca.items()}

    # Evaluar cada candidato contra nombre y dirección
    evaluados = []
    for muni, c in unicos:
        ok_nombre, frac = evaluar_nombre(nom, c)
        if not ok_nombre:
            continue
        lim_via = None
        punto = puntos.get(muni.codigo, (None, None))[1]
        d_via = distancia_a_lineas(c.lat, c.lon, punto.lineas) if punto and punto.lineas else None
        d_pto = punto.distancia(c.lat, c.lon) if punto else None
        # La placa solo cuenta si está sobre la misma vía (o pegada a ella)
        misma_via = coincide_via([c.via], nombres_via(d.tipo, d.numero, d.nombre_via), bool(d.nombre_via)) \
            if d.tipo else False
        placa_ok = bool(d.placa and d.cruce and c.placa and clave(c.placa) == clave(f"{d.cruce}{d.placa}")
                        and (misma_via or (d_via is not None and d_via <= 60)))
        coherente = direccion_coherente(d, c, misma_via, d_via)
        d_cc = min((distancia_m(c.lat, c.lon, m.lat, m.lon) for _, m in centros), default=None)
        cc_ok = d_cc is not None and d_cc <= MAX_DIST_CC
        if placa_ok or coherente or cc_ok or \
                (punto and punto.precision == "EXACTA" and d_pto is not None and d_pto <= RADIO_MISMA_SEDE):
            dir_ok = True
        elif d_via is not None or d_pto is not None:
            sede_exacta = (bool(nom.sucursal) and frac >= 0.99) or \
                (bool(nom.marca) and not nom.sucursal and locales_marca.get(muni.codigo) == 1)
            if sede_exacta:
                lim_via = MAX_DIST_VIA
            elif d_pto is None:    # solo se conoce la vía: el punto debe estar sobre ella
                lim_via = MAX_DIST_SOLO_VIA
            else:
                lim_via = MAX_DIST_VIA_ESTRICTA
            # Un punto geocodificado sobre la placa (Google) es preciso: margen más corto que un cruce
            lim_pto = MAX_DIST_CRUCE
            if punto and punto.precision in ("EXACTA", "INTERPOLADA") and not punto.lineas:
                lim_pto = MAX_DIST_VIA if sede_exacta else MAX_DIST_PLACA
            dir_ok = (d_via is not None and d_via <= lim_via and (d_pto is None or d_pto <= lim_pto)) \
                or (d_via is None and d_pto <= lim_pto)
        else:
            dir_ok = None   # no hay información de dirección verificable
        evaluados.append({"muni": muni, "c": c, "frac": frac, "d_via": d_via, "d_pto": d_pto,
                          "placa_ok": placa_ok, "dir_ok": dir_ok, "lim_via": lim_via,
                          "coherente": coherente, "d_cc": d_cc if cc_ok else None, "misma_via": misma_via})

    hay_info_dir = any(p.lineas or p.lat is not None for _, p in puntos.values()) or bool(centros)
    if d.via_principal and not hay_info_dir and not d.ausente:
        obs.append("La dirección registrada no pudo ubicarse en la cartografía; no se usó para validar.")

    # Palabras del nombre del candidato que no son marca, sede ni genéricas: un nombre limpio
    # ("Éxito Américas") se prefiere a un negocio dentro de la tienda ("Éxito Américas Natural Light")
    propias = {t for a in nom.alias for t in a} | set(nom.sucursal) | PALABRAS_GENERICAS | SEDE_GENERICA | \
        {"WOW", "CC", "COMERCIAL", "SUPERMERCADOS", "ALMACEN", "SUPERALMACEN"} | \
        {t for m in DEPARTAMENTOS for t in m.split()}
    extra = lambda e: len(set(_expandir_abreviaturas(tokens(e["c"].nombre))) - propias)

    elegido, tipo, estado = None, "", ""
    ok = []
    if hay_info_dir:
        ok = [e for e in evaluados if e["dir_ok"]]
        fuerza = lambda e: (e["placa_ok"] or e["coherente"], e["frac"], -extra(e))
        if ok:
            grupos = _agrupar([e["c"] for e in ok])
            if len(grupos) > 1:
                ok.sort(key=lambda e: (-fuerza(e)[0], -e["frac"], extra(e),
                                       e["d_via"] if e["d_via"] is not None else e["d_pto"] or 0))
                mejor = ok[0]
                seg = next(e for e in ok[1:] if distancia_m(e["c"].lat, e["c"].lon, mejor["c"].lat,
                                                            mejor["c"].lon) > RADIO_MISMA_SEDE)
                if fuerza(mejor) > fuerza(seg):
                    elegido = mejor
                else:
                    estado = "REVISAR"
                    obs.append(f"{len(grupos)} sedes compatibles con la dirección: " + " | ".join(
                        f"{g[0].nombre} ({_fmt(g[0].lat, g[0].lon)})" for g in grupos[:4]) + ".")
            else:
                elegido = sorted(ok, key=lambda e: (-e["placa_ok"], -e["coherente"], -e["frac"]))[0]
            if elegido:
                tipo = "SEDE_IDENTIFICADA_DIRECCION_COINCIDENTE" if elegido["placa_ok"] else "SEDE_IDENTIFICADA"
        elif evaluados:
            cerca = min(evaluados, key=lambda e: e["d_pto"] if e["d_pto"] is not None else e["d_via"] or 1e9)
            dist = cerca["d_pto"] if cerca["d_pto"] is not None else cerca["d_via"]
            exacta = any(p.precision in ("EXACTA", "ESQUINA") and cod in principales
                         for cod, (_, p) in puntos.items())
            lejana = dist is None or (cerca["frac"] < 0.5 and (
                dist > MAX_DIST_CANDIDATO or cerca["muni"].codigo not in principales))
        if evaluados and not ok and (exacta or lejana):
            # Las sedes encontradas de la marca no son esta: no se proponen; se recurre a la dirección
            obs.append(f"Las sedes de la marca encontradas no corresponden a la dirección (la más cercana, "
                       f"'{cerca['c'].nombre}', está a {dist or 0:.0f} m).")
        elif evaluados and not ok:
            distancias = []
            if cerca["d_via"] is not None:
                distancias.append(f"{cerca['d_via']:.0f} m de la vía principal (máx. {cerca['lim_via']})")
            if cerca["d_pto"] is not None:
                distancias.append(f"{cerca['d_pto']:.0f} m del cruce de la dirección (máx. {MAX_DIST_CRUCE})")
            estado = "REVISAR"
            obs.append(f"Se encontró '{cerca['c'].nombre}' pero no es compatible con la dirección registrada: "
                       f"{', '.join(distancias) or 'sin distancia'}; nombre de sede coincide al "
                       f"{cerca['frac']:.0%}. Candidato {_fmt(cerca['c'].lat, cerca['c'].lon)} ({cerca['c'].place_id}).")
            res.update(NOMBRE_ENCONTRADO=cerca["c"].nombre, DIRECCION_ENCONTRADA=cerca["c"].direccion,
                       CIUDAD_ENCONTRADA=cerca["muni"].nombre, PLACE_ID=cerca["c"].place_id or "",
                       FUENTE_COORDENADAS=cerca["c"].fuente, TIPO_COINCIDENCIA="DISCREPANCIA_DIRECCION")
    else:
        if nom.marca:
            fuertes = [e for e in evaluados if e["frac"] >= (0.5 if nom.sucursal else 1.0)]
        else:
            fuertes = evaluados
        grupos = _agrupar([e["c"] for e in fuertes])
        mejor_por_grupo = []
        for g_ in grupos:
            en_grupo = [e for e in fuertes if e["c"] in g_]
            mejor_por_grupo.append(max(en_grupo, key=lambda e: (e["frac"], e["coherente"], -extra(e))))
        mejor_por_grupo.sort(key=lambda e: (e["frac"], e["coherente"], -extra(e)), reverse=True)
        clave_m = lambda e: (e["frac"], e["coherente"], -extra(e))
        if len(grupos) == 1 or (len(grupos) > 1 and clave_m(mejor_por_grupo[0]) > clave_m(mejor_por_grupo[1])):
            elegido = mejor_por_grupo[0]
            tipo = "SEDE_IDENTIFICADA"
            obs.append("Sede identificada por nombre y municipio; la dirección no permitió verificarla" +
                       (f" (se descartaron {len(grupos) - 1} sedes con nombre menos preciso)." if len(grupos) > 1 else "."))
        elif len(grupos) > 1:
            estado = "REVISAR"
            obs.append(f"{len(grupos)} sedes posibles con ese nombre: " + " | ".join(
                f"{g[0].nombre} ({_fmt(g[0].lat, g[0].lon)})" for g in grupos[:4]) + ".")
        elif evaluados:
            estado = "REVISAR"
            grupos = _agrupar([e["c"] for e in evaluados])
            obs.append(f"Hay {len(grupos)} sede(s) de la marca en la zona, pero ninguna coincide con el nombre "
                       "de la sede: " + " | ".join(f"{g[0].nombre} ({_fmt(g[0].lat, g[0].lon)}, {g[0].place_id})"
                                                  for g in grupos[:4]) + ".")

    if elegido:
        c, muni = elegido["c"], elegido["muni"]
        res.update(NOMBRE_ENCONTRADO=c.nombre, DIRECCION_ENCONTRADA=c.direccion,
                   CIUDAD_ENCONTRADA=muni.nombre, FUENTE_COORDENADAS=c.fuente,
                   PLACE_ID=c.place_id or "", TIPO_COINCIDENCIA=tipo)
        if elegido["d_via"] is not None or elegido["d_pto"] is not None or elegido["d_cc"] is not None:
            partes = []
            if elegido["placa_ok"]:
                partes.append(f"placa {c.placa} coincide")
            if elegido["coherente"]:
                partes.append(f"ubicada sobre {c.via}" + (f" {c.placa}" if c.placa else "") +
                              ", coherente con la nomenclatura de la dirección")
            if elegido["d_cc"] is not None:
                cc = min(centros, key=lambda x: distancia_m(c.lat, c.lon, x[1].lat, x[1].lon))[1]
                partes.append(f"{elegido['d_cc']:.0f} m del centro comercial '{cc.nombre}'")
            if elegido["misma_via"] and not elegido["coherente"]:
                partes.append(f"figura sobre {c.via}" + (f" (placa {c.placa} en la fuente)" if c.placa else ""))
            if elegido["d_via"] is not None and not (elegido["misma_via"] and elegido["d_via"] > MAX_DIST_VIA):
                partes.append(f"{elegido['d_via']:.0f} m de la vía principal")
            if elegido["d_pto"] is not None:
                partes.append(f"{elegido['d_pto']:.0f} m del punto de la dirección")
            obs.append("Validación de dirección: " + ", ".join(partes) + ".")
        if c.placa and d.cruce and d.placa and not elegido["placa_ok"] and \
                (elegido["misma_via"] or elegido["coherente"]) and clave(c.placa) != clave(f"{d.cruce}{d.placa}"):
            obs.append(f"La placa en la fuente ({c.placa}) difiere de la registrada ({d.cruce}-{d.placa}); "
                       "revisar la dirección de la base.")
        fallos = validar_coordenadas(c.lat, c.lon, c.pais or "co", muni)
        estado = "GEORREFERENCIADO"
        if fallos:
            estado = "REVISAR"
            obs.append("Validación de coordenadas: " + "; ".join(fallos) + ".")
        # El centro comercial citado en la propia dirección ("CC PZ AMERICAS") respalda la sede aunque
        # la parte numérica de la dirección se ubique lejos
        toks_compl = set(_expandir_abreviaturas(tokens(d.complemento)))
        cc_en_direccion = elegido["d_cc"] is not None and \
            any(t in toks_compl for t in nom.sucursal if t not in SEDE_GENERICA)
        if cc_en_direccion and elegido["d_pto"] is not None and elegido["d_pto"] > MAX_DIST_CONFLICTO:
            obs.append(f"La parte numérica de la dirección se ubica a {elegido['d_pto'] / 1000:.1f} km, pero la "
                       "dirección cita el mismo centro comercial; revisar la nomenclatura registrada.")
        elif elegido["d_pto"] is not None and elegido["d_pto"] > MAX_DIST_CONFLICTO and not elegido["placa_ok"]:
            estado = "REVISAR"
            tipo = "DISCREPANCIA_DIRECCION"
            res["TIPO_COINCIDENCIA"] = tipo
            obs.append(f"DISCREPANCIA: la sede se identificó por nombre, pero la dirección registrada se ubica a "
                       f"{elegido['d_pto'] / 1000:.1f} km; confirmar cuál es la correcta.")
        if muni.codigo != ciudad_cod:
            # Fuera de la ciudad registrada se exige evidencia fuerte de dirección
            dir_fuerte = elegido["placa_ok"] or elegido["coherente"] or \
                (elegido["d_via"] is not None and elegido["d_via"] <= 100) or \
                (elegido["d_pto"] is not None and elegido["d_pto"] <= 200)
            if muni.codigo == nom.municipio_en_nombre or ((hay_info_dir or elegido["coherente"]) and dir_fuerte):
                obs.append(f"ATENCIÓN: CIUDAD registrada '{ciudad}' no coincide; la sede está en "
                           f"{muni.nombre} (respaldo: " +
                           ("municipio indicado en el nombre" if muni.codigo == nom.municipio_en_nombre
                            else f"nomenclatura coherente: {c.via} {c.placa}".strip()
                            if elegido["coherente"] and not hay_info_dir
                            else "dirección verificada en ese municipio") + "). Corregir CIUDAD en la base.")
            else:
                estado = "REVISAR"
                obs.append(f"Discrepancia de municipio: registrada '{ciudad}', sede probable en {muni.nombre}, "
                           "sin evidencia de dirección suficiente para confirmarla.")
        if estado == "GEORREFERENCIADO":
            res.update(LATITUD_VALIDADA=round(c.lat, 7), LONGITUD_VALIDADA=round(c.lon, 7))
        else:
            obs.append(f"Candidato propuesto: {_fmt(c.lat, c.lon)}.")
    elif not estado:
        # Sin establecimiento: usar la dirección solo si es exacta
        exactos = [(m, p) for cod, (m, p) in puntos.items()
                   if p.precision in ("EXACTA", "ESQUINA") and p.lat is not None and cod in principales]
        aprox = [(m, p) for cod, (m, p) in puntos.items()
                 if p.precision in ("CRUCE_APROX", "CRUCE_AMBIGUO", "INTERPOLADA") and p.lat is not None
                 and cod in principales]
        if exactos:
            muni, p = exactos[0]
            fallos = validar_coordenadas(p.lat, p.lon, "co", muni)
            estado = "REVISAR" if fallos else "GEORREFERENCIADO"
            tipo = "DIRECCION_COINCIDENTE"
            res.update(DIRECCION_ENCONTRADA=p.descripcion, CIUDAD_ENCONTRADA=muni.nombre,
                       FUENTE_COORDENADAS=p.fuente, PLACE_ID=p.place_id or "", TIPO_COINCIDENCIA=tipo)
            obs.append("Ubicación de la dirección (" + ("esquina" if p.precision == "ESQUINA" else "placa") +
                       "); el establecimiento no figura en la fuente, no es una sede verificada.")
            if fallos:
                obs.append("; ".join(fallos) + f". Candidato: {_fmt(p.lat, p.lon)}.")
            else:
                res.update(LATITUD_VALIDADA=round(p.lat, 7), LONGITUD_VALIDADA=round(p.lon, 7))
        elif centros:
            muni, cc = centros[0]
            estado = "REVISAR"
            res.update(NOMBRE_ENCONTRADO=cc.nombre, DIRECCION_ENCONTRADA=cc.direccion,
                       CIUDAD_ENCONTRADA=muni.nombre, FUENTE_COORDENADAS=cc.fuente,
                       PLACE_ID=cc.place_id or "", TIPO_COINCIDENCIA="UBICACION_APROXIMADA")
            obs.append(f"UBICACIÓN APROXIMADA (no verificada): centro comercial '{cc.nombre}' "
                       f"({_fmt(cc.lat, cc.lon)}); el local no figura en la fuente.")
        elif aprox:
            muni, p = aprox[0]
            estado = "REVISAR"
            res.update(DIRECCION_ENCONTRADA=p.descripcion, CIUDAD_ENCONTRADA=muni.nombre,
                       FUENTE_COORDENADAS=p.fuente, TIPO_COINCIDENCIA="UBICACION_APROXIMADA")
            obs.append(f"UBICACIÓN APROXIMADA (no verificada): {p.descripcion}, {_fmt(p.lat, p.lon)}; "
                       "el establecimiento no se encontró en la fuente.")
        else:
            estado = "SIN_RESULTADO"
            if d.ausente:
                obs.append("Sin dirección y sin coincidencia confiable por nombre y ciudad.")
            else:
                obs.append("No se encontró el establecimiento ni se pudo ubicar la dirección.")
    if nom.ciudades_marca and ciudad_cod not in nom.ciudades_marca and \
            res["CIUDAD_ENCONTRADA"] not in [NOMBRE_OFICIAL.get(x, x.title()) for x in nom.ciudades_marca] \
            and estado != "GEORREFERENCIADO":
        obs.append(f"{nom.marca_busqueda} opera principalmente en "
                   f"{', '.join(x.title() for x in nom.ciudades_marca)}; verificar la ciudad registrada.")
    res["ESTADO_GEORREFERENCIACION"] = estado
    res["OBSERVACIONES"] = " ".join(obs)
    res["CONSULTAS"] = consultas
    return res


# ---------------------------------------------------------------------------
# Correcciones validadas por una persona (correcciones_manuales.xlsx)
# ---------------------------------------------------------------------------
COLUMNAS_CORRECCION = ["SAP", "PDV", "CIUDAD", "DIRECCION", "ESTADO_ACTUAL", "MOTIVO",
                       "BUSCAR_EN_GOOGLE_MAPS", "VER_CANDIDATO", "ENLACE_O_COORDENADAS",
                       "CIUDAD_CORRECTA", "FUENTE", "NOTA"]
COLUMNAS_EDITABLES = ["ENLACE_O_COORDENADAS", "CIUDAD_CORRECTA", "FUENTE", "NOTA"]
RE_COORDENADAS = [
    ("pin", re.compile(r"!3d(-?\d{1,2}\.\d+)!4d(-?\d{1,3}\.\d+)")),           # pin de un lugar
    ("consulta", re.compile(r"[?&](?:q|query|ll|destination)=(-?\d{1,2}\.\d+),\s*(-?\d{1,3}\.\d+)")),
    ("vista", re.compile(r"@(-?\d{1,2}\.\d+),(-?\d{1,3}\.\d+)")),             # centro del mapa
    ("texto", re.compile(r"^\s*(-?\d{1,2}[.,]\d+)\s*[,; ]\s*(-?\d{1,3}[.,]\d+)\s*$")),
]


def extraer_coordenadas(texto, resolver_enlaces: bool = True) -> tuple[float, float, str] | None:
    """Lee coordenadas de un enlace de Google Maps/OSM o de un texto '4.82, -75.68'.
    Devuelve (lat, lon, origen) o None."""
    if es_vacio(texto):
        return None
    texto = str(texto).strip()
    if resolver_enlaces and re.match(r"https?://(maps\.app\.goo\.gl|goo\.gl/maps)/", texto):
        try:  # enlaces cortos de "Compartir": se sigue la redirección para leer la URL completa
            texto = requests.get(texto, allow_redirects=True, timeout=20,
                                 headers={"User-Agent": "Mozilla/5.0"}).url
        except requests.RequestException as e:
            log.warning("No se pudo abrir el enlace corto %s: %s", texto, e)
            return None
    for origen, patron in RE_COORDENADAS:
        m = patron.search(texto)
        if m:
            return float(m.group(1).replace(",", ".")), float(m.group(2).replace(",", ".")), origen
    return None


def cargar_correcciones(ruta: Path) -> dict[str, dict]:
    """Filas de la plantilla en las que alguien llenó ENLACE_O_COORDENADAS."""
    if not ruta.exists():
        return {}
    df = pd.read_excel(ruta, dtype=object)
    salida = {}
    for f in df.to_dict("records"):
        if not es_vacio(f.get("SAP")) and not es_vacio(f.get("ENLACE_O_COORDENADAS")):
            salida[str(f["SAP"]).split(".")[0]] = f
    return salida


def aplicar_correccion(resultado: dict, corr: dict, fila: dict, prov) -> dict:
    """Valida la ubicación indicada por una persona y, si es coherente, la adopta."""
    r = dict(resultado)
    previo = f"[Automático: {resultado.get('ESTADO_GEORREFERENCIACION')}] {resultado.get('OBSERVACIONES', '')}"
    coords = extraer_coordenadas(corr.get("ENLACE_O_COORDENADAS"))
    if not coords:
        r["OBSERVACIONES"] = ("[Corrección manual ignorada: no se reconocen coordenadas en "
                              f"'{corr.get('ENLACE_O_COORDENADAS')}'] " + resultado.get("OBSERVACIONES", ""))
        return r
    lat, lon, origen = coords
    ciudad = corr.get("CIUDAD_CORRECTA") if not es_vacio(corr.get("CIUDAD_CORRECTA")) else fila["CIUDAD"]
    muni = prov.municipio(codigo_municipio(ciudad))
    fallos = validar_coordenadas(lat, lon, "co")
    if muni and not muni.contiene(lat, lon):
        fallos.append(f"el punto no está dentro de {muni.nombre}")
    if not muni:
        fallos.append(f"no se pudo validar el municipio '{ciudad}'")
    obs = [f"[Corrección manual] Coordenadas {_fmt(lat, lon)} tomadas "
           + {"pin": "del pin del enlace", "consulta": "del enlace", "texto": "del texto ingresado",
              "vista": "del centro del mapa del enlace (verificar que el local esté centrado)"}[origen] + "."]
    if not es_vacio(corr.get("FUENTE")):
        obs.append(f"Fuente: {corr['FUENTE']}.")
    if not es_vacio(corr.get("NOTA")):
        obs.append(f"Nota: {corr['NOTA']}.")
    lat_auto = resultado.get("LATITUD_VALIDADA")
    m = re.search(r"(-?\d{1,2}\.\d{5,}), (-?\d{1,3}\.\d{5,})", resultado.get("OBSERVACIONES", ""))
    ref = (lat_auto, resultado.get("LONGITUD_VALIDADA")) if lat_auto is not None else \
        ((float(m.group(1)), float(m.group(2))) if m else None)
    if ref:
        obs.append(f"Está a {distancia_m(lat, lon, ref[0], ref[1]):.0f} m de la ubicación automática.")
    if muni and codigo_municipio(fila["CIUDAD"]) != muni.codigo:
        obs.append(f"Municipio corregido: {fila['CIUDAD']} -> {muni.nombre}.")
    if fallos:
        r.update(ESTADO_GEORREFERENCIACION="REVISAR", LATITUD_VALIDADA=None, LONGITUD_VALIDADA=None)
        obs.append("No se adoptó: " + "; ".join(fallos) + ".")
    else:
        r.update(ESTADO_GEORREFERENCIACION="GEORREFERENCIADO", LATITUD_VALIDADA=round(lat, 7),
                 LONGITUD_VALIDADA=round(lon, 7), CIUDAD_ENCONTRADA=muni.nombre,
                 TIPO_COINCIDENCIA="VALIDACION_MANUAL", PLACE_ID="",
                 FUENTE_COORDENADAS="Validación manual" + (f" ({corr['FUENTE']})" if not es_vacio(corr.get("FUENTE")) else ""),
                 NOMBRE_ENCONTRADO=str(fila["PDV"]), DIRECCION_ENCONTRADA="")
    r["OBSERVACIONES"] = " ".join(obs) + " " + previo
    return r


def url_busqueda_maps(texto: str) -> str:
    from urllib.parse import quote
    return "https://www.google.com/maps/search/?api=1&query=" + quote(texto)


def generar_plantilla_correcciones(salida: pd.DataFrame, ruta: Path) -> int:
    """Crea/actualiza la plantilla con los registros no georreferenciados automáticamente.
    Conserva lo que ya se haya escrito en las columnas editables."""
    previas = {}
    if ruta.exists():
        for f in pd.read_excel(ruta, dtype=object).to_dict("records"):
            previas[str(f.get("SAP")).split(".")[0]] = f
    filas = []
    for f in salida.to_dict("records"):
        sap = str(f["SAP"])
        manual = f.get("TIPO_COINCIDENCIA") == "VALIDACION_MANUAL"
        escrito = any(not es_vacio(previas.get(sap, {}).get(c)) for c in COLUMNAS_EDITABLES)
        if f["ESTADO_GEORREFERENCIACION"] == "GEORREFERENCIADO" and not manual and not escrito:
            continue
        nom = analizar_nombre(f["PDV"], f["SAP"])
        d = parsear_direccion(f["DIRECCION"])
        consulta = ", ".join(x for x in [nom.nombre_limpio(), d.normalizada(), str(f["CIUDAD"]).title(), "Colombia"] if x)
        m = re.search(r"(-?\d{1,2}\.\d{5,}), (-?\d{1,3}\.\d{5,})", str(f.get("OBSERVACIONES") or ""))
        motivo = re.sub(r"\[Corrección manual\].*?\[Automático: \w+\] ", "", str(f.get("OBSERVACIONES") or ""))
        motivo = re.sub(r"^Original: [^.]*\.( [^.]*Colombia\.)? ?", "", motivo)
        fila = {"SAP": f["SAP"], "PDV": f["PDV"], "CIUDAD": f["CIUDAD"], "DIRECCION": f["DIRECCION"],
                "ESTADO_ACTUAL": f["ESTADO_GEORREFERENCIACION"], "MOTIVO": motivo[:500],
                "BUSCAR_EN_GOOGLE_MAPS": url_busqueda_maps(consulta),
                "VER_CANDIDATO": url_busqueda_maps(f"{m.group(1)},{m.group(2)}") if m else ""}
        for col in COLUMNAS_EDITABLES:
            v = previas.get(sap, {}).get(col)
            fila[col] = None if es_vacio(v) else v
        filas.append(fila)
    wb = Workbook()
    ws = wb.active
    ws.title = "Correcciones"
    ws.append(COLUMNAS_CORRECCION)
    editable = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = editable if c.value in COLUMNAS_EDITABLES else PatternFill("solid", fgColor="DDEBF7")
    for fila in filas:
        ws.append([fila[c] for c in COLUMNAS_CORRECCION])
        for col in ("BUSCAR_EN_GOOGLE_MAPS", "VER_CANDIDATO"):
            celda = ws.cell(ws.max_row, COLUMNAS_CORRECCION.index(col) + 1)
            if celda.value:
                celda.hyperlink, celda.value = celda.value, "Abrir mapa"
                celda.font = Font(color="0563C1", underline="single")
        ws.cell(ws.max_row, COLUMNAS_CORRECCION.index("MOTIVO") + 1).alignment = Alignment(wrap_text=True, vertical="top")
    for i, ancho in enumerate([11, 40, 14, 28, 18, 70, 16, 14, 45, 16, 22, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "C2"
    ayuda = wb.create_sheet("Instrucciones")
    for linea in [
        "Cómo completar esta plantilla",
        "1. Abre 'BUSCAR_EN_GOOGLE_MAPS' (o 'VER_CANDIDATO' si hay una ubicación propuesta) y localiza la sede exacta.",
        "2. Haz clic derecho sobre el local -> copia las coordenadas (p. ej. 4.823696, -75.684668), "
        "o usa Compartir -> Copiar enlace.",
        "3. Pega las coordenadas o el enlace en la columna amarilla ENLACE_O_COORDENADAS.",
        "4. Si el municipio real es otro, escríbelo en CIUDAD_CORRECTA (p. ej. Manizales).",
        "5. Indica en FUENTE cómo lo verificaste (Google Maps, visita, llamada a la tienda...).",
        "6. Guarda el archivo y ejecuta: uv run georreferenciar_pdv.py",
        "El script valida que el punto esté en Colombia y dentro del municipio antes de adoptarlo; "
        "si no, el registro queda en REVISAR con el motivo.",
        "Puedes agregar filas de registros ya georreferenciados para corregirlos (solo SAP + coordenadas).",
    ]:
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True)
    ayuda.column_dimensions["A"].width = 120
    wb.save(ruta)
    return len(filas)


# ---------------------------------------------------------------------------
# Progreso, exportación y resumen
# ---------------------------------------------------------------------------
def firma_registro(fila: dict, proveedor: str) -> str:
    datos = [str(fila[c]) for c in ("SAP", "PDV", "CIUDAD", "DIRECCION", "LATITUD", "LONGIUTD")]
    return Cache.llave(datos, proveedor, VERSION_ALGORITMO)


def cargar_progreso(ruta: Path) -> dict:
    return json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}


def guardar_progreso(ruta: Path, progreso: dict) -> None:
    escribir_json(ruta, progreso)


def construir_salida(df: pd.DataFrame, resultados: dict, actualizar_originales: bool = True) -> pd.DataFrame:
    salida = df[COLUMNAS_ORIGINALES].copy()
    salida["LATITUD_ORIGINAL"] = df["LATITUD"]
    salida["LONGITUD_ORIGINAL"] = df["LONGIUTD"]
    for col in COLUMNAS_NUEVAS[2:]:
        salida[col] = [resultados.get(str(s), {}).get(col) for s in df["SAP"]]
    if actualizar_originales:
        ok = salida["ESTADO_GEORREFERENCIACION"] == "GEORREFERENCIADO"
        salida["LATITUD"] = salida["LATITUD"].astype(object)
        salida["LONGIUTD"] = salida["LONGIUTD"].astype(object)
        salida.loc[ok, "LATITUD"] = salida.loc[ok, "LATITUD_VALIDADA"]
        salida.loc[ok, "LONGIUTD"] = salida.loc[ok, "LONGITUD_VALIDADA"]
    return salida[COLUMNAS_ORIGINALES + COLUMNAS_NUEVAS]


def generar_resumen(salida: pd.DataFrame, proveedor: str, modo: str) -> list[tuple[str, object]]:
    estados = salida["ESTADO_GEORREFERENCIACION"].fillna("NO_PROCESADO")
    validas_antes = sum(diagnosticar_coordenadas(a, b)[0] == "VALIDA"
                        for a, b in zip(salida["LATITUD_ORIGINAL"], salida["LONGITUD_ORIGINAL"]))
    return [
        ("Total de puntos de venta", len(salida)),
        ("Coordenadas válidas antes del procesamiento", int(validas_antes)),
        ("Registros georreferenciados", int((estados == "GEORREFERENCIADO").sum())),
        ("   de ellos, por validación manual", int(((estados == "GEORREFERENCIADO") &
                                                   (salida["TIPO_COINCIDENCIA"] == "VALIDACION_MANUAL")).sum())),
        ("Registros pendientes de revisión", int((estados == "REVISAR").sum())),
        ("Registros sin resultados", int((estados == "SIN_RESULTADO").sum())),
        ("Errores de consulta", int((estados == "ERROR").sum())),
        ("Registros no procesados", int((estados == "NO_PROCESADO").sum())),
    ], [
        ("Fecha de ejecución", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Proveedor", proveedor),
        ("Modo", modo),
        ("Versión del algoritmo", VERSION_ALGORITMO),
    ]


COLORES_ESTADO = {"GEORREFERENCIADO": "C6EFCE", "REVISAR": "FFEB9C",
                  "SIN_RESULTADO": "F4CCCC", "ERROR": "E06666"}


def exportar_excel(salida: pd.DataFrame, ruta: Path, proveedor: str, modo: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA_ENTRADA
    negrita = Font(bold=True)
    encabezado = PatternFill("solid", fgColor="DDEBF7")
    ws.append(list(salida.columns))
    for c in ws[1]:
        c.font, c.fill = negrita, encabezado
    col_estado = list(salida.columns).index("ESTADO_GEORREFERENCIACION") + 1
    for fila in salida.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and math.isnan(v)) or v is pd.NaT else
                   (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in fila])
        estado = ws.cell(ws.max_row, col_estado)
        if estado.value in COLORES_ESTADO:
            estado.fill = PatternFill("solid", fgColor=COLORES_ESTADO[estado.value])
    anchos = {"PDV": 42, "DIRECCION": 30, "NOMBRE_ENCONTRADO": 32, "DIRECCION_ENCONTRADA": 55,
              "OBSERVACIONES": 90, "ESTADO_GEORREFERENCIACION": 20, "TIPO_COINCIDENCIA": 26,
              "FUENTE_COORDENADAS": 24, "PLACE_ID": 20}
    for i, col in enumerate(salida.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = anchos.get(col, 14)
        if col == "fech":
            for celda in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in celda:
                    c.number_format = "yyyy-mm-dd"
        if col in ("LATITUD", "LONGIUTD", "LATITUD_VALIDADA", "LONGITUD_VALIDADA"):
            for celda in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in celda:
                    if isinstance(c.value, float):
                        c.number_format = "0.0000000"
        if col == "OBSERVACIONES":
            for celda in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in celda:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    totales, meta = generar_resumen(salida, proveedor, modo)
    rs = wb.create_sheet("Resumen")
    rs.append(["Indicador", "Valor"])
    for c in rs[1]:
        c.font, c.fill = negrita, encabezado
    for fila in totales:
        rs.append(list(fila))
    rs.append([])
    for fila in meta:
        rs.append(list(fila))
    rs.append([])
    rs.append(["Registros no georreferenciados y motivo"])
    rs.cell(rs.max_row, 1).font = negrita
    rs.append(["SAP", "PDV", "CIUDAD", "DIRECCION", "ESTADO", "OBSERVACIONES"])
    for c in rs[rs.max_row]:
        c.font, c.fill = negrita, encabezado
    pendientes = salida[salida["ESTADO_GEORREFERENCIACION"] != "GEORREFERENCIADO"]
    for f in pendientes.itertuples(index=False):
        rs.append([f.SAP, f.PDV, f.CIUDAD, f.DIRECCION, f.ESTADO_GEORREFERENCIACION, f.OBSERVACIONES])
        rs.cell(rs.max_row, 6).alignment = Alignment(wrap_text=True, vertical="top")
    for col, ancho in zip("ABCDEF", (44, 42, 16, 30, 18, 100)):
        rs.column_dimensions[col].width = ancho
    wb.save(ruta)


# ---------------------------------------------------------------------------
# Orquestación
# ---------------------------------------------------------------------------
def crear_proveedor(nombre: str, cache: Cache):
    osm = Nominatim(cache)
    key = os.environ.get("GOOGLE_MAPS_API_KEY", "").strip()
    if nombre == "auto":
        nombre = "google" if key else "osm"
    if nombre == "google":
        if not key:
            raise SystemExit("Falta la variable de entorno GOOGLE_MAPS_API_KEY para usar --proveedor google.")
        return "google", GooglePlaces(cache, key, osm)
    return "osm", osm


def imprimir_diagnostico(df: pd.DataFrame) -> None:
    print(f"{len(df)} registros en {HOJA_ENTRADA}\n")
    conteo: dict[str, int] = {}
    for f in df.to_dict("records"):
        diag, _ = diagnosticar_coordenadas(f["LATITUD"], f["LONGIUTD"])
        conteo[diag] = conteo.get(diag, 0) + 1
        d = parsear_direccion(f["DIRECCION"])
        nom = analizar_nombre(f["PDV"], f["SAP"])
        print(f"{f['SAP']} | {diag:<9} | {f['PDV'][:40]:<40} | {str(f['CIUDAD'])[:13]:<13} | "
              f"{'(SIN DIRECCION)' if d.ausente else d.normalizada():<40} | marca={nom.marca or '-'} "
              f"sede={' '.join(nom.sucursal) or '-'} muni_nombre={nom.municipio_en_nombre or '-'}"
              + (f" | {'; '.join(d.problemas)}" if d.problemas and not d.ausente else ""))
    print("\nDiagnóstico de coordenadas originales:", conteo)
    print("Sin dirección:", sum(direccion_ausente(x) for x in df["DIRECCION"]))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--entrada", type=Path, default=ARCHIVO_ENTRADA)
    ap.add_argument("--salida", type=Path)
    ap.add_argument("--proveedor", choices=["auto", "osm", "google"], default="auto")
    ap.add_argument("--correcciones", type=Path, default=ARCHIVO_CORRECCIONES,
                    help="Excel con ubicaciones validadas por una persona")
    ap.add_argument("--prueba", action="store_true", help="Procesa solo los primeros 5 registros")
    ap.add_argument("--limite", type=int, help="Procesa solo los primeros N registros")
    ap.add_argument("--sap", nargs="*", help="Procesa solo estos códigos SAP")
    ap.add_argument("--reprocesar", action="store_true", help="Ignora el progreso guardado")
    ap.add_argument("--diagnostico", action="store_true", help="Solo analiza la base (sin red)")
    ap.add_argument("--no-actualizar-originales", action="store_true",
                    help="No escribe las coordenadas validadas en LATITUD/LONGIUTD")
    args = ap.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    DIR_TRABAJO.mkdir(exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.StreamHandler(sys.stdout),
                                  logging.FileHandler(DIR_TRABAJO / "ejecucion.log", encoding="utf-8")])
    df = leer_excel(args.entrada)
    if args.diagnostico:
        imprimir_diagnostico(df)
        return 0

    limite = 5 if args.prueba else args.limite
    seleccion = df
    if args.sap:
        seleccion = df[df["SAP"].astype(str).isin(args.sap)]
    if limite:
        seleccion = seleccion.head(limite)
    modo = "prueba" if (limite or args.sap) else "completo"
    salida_ruta = args.salida or (ARCHIVO_SALIDA_PRUEBA if modo == "prueba" else ARCHIVO_SALIDA)
    if salida_ruta.resolve() == args.entrada.resolve():
        raise SystemExit("La salida no puede sobrescribir el archivo original.")

    cache = Cache(DIR_TRABAJO / "cache_consultas.json")
    nombre_prov, prov = crear_proveedor(args.proveedor, cache)
    ruta_progreso = DIR_TRABAJO / f"progreso_{nombre_prov}.json"
    progreso = {} if args.reprocesar else cargar_progreso(ruta_progreso)
    log.info("Proveedor: %s | registros a procesar: %d | modo: %s", nombre_prov, len(seleccion), modo)

    for fila in seleccion.to_dict("records"):
        sap = str(fila["SAP"])
        firma = firma_registro(fila, nombre_prov)
        previo = progreso.get(sap)
        if previo and previo["firma"] == firma and previo["resultado"]["ESTADO_GEORREFERENCIACION"] != "ERROR":
            log.info("%s %s: ya procesado (%s)", sap, fila["PDV"], previo["resultado"]["ESTADO_GEORREFERENCIACION"])
            continue
        try:
            resultado = procesar_registro(fila, prov)
        except ErrorServicio as e:
            resultado = {"ESTADO_GEORREFERENCIACION": "ERROR", "OBSERVACIONES": f"Fallo del servicio: {e}",
                         "FUENTE_COORDENADAS": nombre_prov}
        progreso[sap] = {"firma": firma, "resultado": resultado,
                         "procesado": datetime.now(timezone.utc).isoformat()}
        guardar_progreso(ruta_progreso, progreso)
        log.info("%s %s: %s %s", sap, fila["PDV"], resultado["ESTADO_GEORREFERENCIACION"],
                 resultado.get("TIPO_COINCIDENCIA", ""))

    resultados = {s: progreso[s]["resultado"] for s in seleccion["SAP"].astype(str) if s in progreso}
    # Las correcciones manuales se aplican encima del resultado automático en cada ejecución
    correcciones = cargar_correcciones(args.correcciones)
    filas_por_sap = {str(f["SAP"]): f for f in seleccion.to_dict("records")}
    for sap, corr in correcciones.items():
        if sap in resultados:
            resultados[sap] = aplicar_correccion(resultados[sap], corr, filas_por_sap[sap], prov)
            log.info("%s: corrección manual -> %s", sap, resultados[sap]["ESTADO_GEORREFERENCIACION"])
        elif sap not in set(df["SAP"].astype(str)):
            log.warning("Corrección para SAP %s ignorada: no existe en la base", sap)
    salida = construir_salida(seleccion if modo == "prueba" else df, resultados,
                              not args.no_actualizar_originales)
    try:
        exportar_excel(salida, salida_ruta, nombre_prov, modo)
    except PermissionError:
        alterna = salida_ruta.with_name(f"{salida_ruta.stem}_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        log.warning("%s está abierto en otro programa; se guarda como %s", salida_ruta.name, alterna.name)
        salida_ruta = alterna
        exportar_excel(salida, salida_ruta, nombre_prov, modo)
    totales, _ = generar_resumen(salida, nombre_prov, modo)
    log.info("Excel generado: %s", salida_ruta)
    if modo == "completo":
        n = generar_plantilla_correcciones(salida, args.correcciones)
        log.info("Plantilla de correcciones (%d registros por validar): %s", n, args.correcciones)
    for k, v in totales:
        log.info("  %s: %s", k, v)
    log.info("Consultas reales a la red en esta ejecución: %d", prov.http.consultas_reales +
             (prov.osm.http.consultas_reales if hasattr(prov, "osm") else 0))
    return 0


if __name__ == "__main__":
    sys.exit(main())
